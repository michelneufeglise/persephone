"""
Spreadsheet parsing for the Documents panel.

Loads every sheet of an .xlsx/.xlsm (openpyxl), .xls (xlrd), .ods (pandas +
odfpy) or .csv/.tsv file into a pandas DataFrame with:

* header detection — the first of the first 10 rows whose non-empty text
  cells cover ≥ 50% of the columns (and hold ≥ 2 distinct values, so a merged
  title banner is not taken for a header) becomes the header; otherwise the
  columns are named "Column A", "Column B", … after their sheet letters;
* fully-empty rows/columns dropped, merged ranges forward-filled;
* formulas read as their cached values (for .xlsx the formula columns are
  recorded; a formula without a cached value — e.g. a workbook written by a
  script — is evaluated for simple arithmetic / SUM / AVERAGE / MIN / MAX /
  COUNT over same-sheet cells);
* dates as ISO yyyy-mm-dd strings, numbers kept numeric;
* hidden sheets parsed but flagged: they are left out of the document text,
  the meta carries only {name, hidden, rows, cols} for them, the preview
  serves them only with include_hidden, and table queries never touch them.

``extract()`` returns (page_texts, meta) for idp_engine — one page per visible
sheet with a markdown table (≤ 300 rows) and per-column stats — and caches the
parsed frames in a pickle next to the raw file for fast querying
(``load_sheets()``; regenerated when missing or stale).
"""

from __future__ import annotations

import csv
import datetime as _dt
import logging
import math
import pickle
import re
import threading
from pathlib import Path
from typing import Any, Iterable, Optional

log = logging.getLogger("sheets")

SHEET_EXTS = (".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".tsv")
TEXT_ROW_CAP = 300          # rows per sheet written into the document text
HEADER_SCAN_ROWS = 10
CACHE_NAME = "_sheets.pkl"
CACHE_VERSION = 2          # 2: single-cell header rule, hidden-sheet meta
FORMULA_SCAN_MAX_BYTES = 10 * 1024 * 1024   # second openpyxl pass only for files < 10 MB
MERGED_CELLS_MAX_BYTES = 40 * 1024 * 1024   # bigger .xlsx → read-only mode, no merge fill
META_MAX_COLUMNS = 200


def is_sheet_filename(filename: str) -> bool:
    return (filename or "").lower().endswith(SHEET_EXTS)


# ── Small value helpers ──────────────────────────────────────────────────────

def col_letter(idx: int) -> str:
    """0-based column index → spreadsheet letters (0 → A, 26 → AA)."""
    idx += 1
    out = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        out = chr(65 + rem) + out
    return out


def _letter_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and not (
        isinstance(v, float) and math.isnan(v)
    )


def _is_empty(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, str) and not v.strip():
        return True
    return False


_NUM_RE = re.compile(r"^[-+]?(?:\d+(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d+)?$")
_THOUSANDS_RE = re.compile(r"^[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?$")
_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}(?::\d{2})?)?$")


def coerce_text_number(s: str) -> Any:
    """'1234' → 1234, '1,234.5' → 1234.5, anything else unchanged."""
    t = s.strip()
    if _NUM_RE.match(t):
        try:
            f = float(t)
        except ValueError:
            return s
        if re.match(r"^[-+]?\d+$", t) and abs(f) < 2**53:
            return int(t)
        return f
    if _THOUSANDS_RE.match(t):
        try:
            f = float(t.replace(",", ""))
            return int(f) if f.is_integer() and "." not in t else f
        except ValueError:
            return s
    return s


def iso_value(v: Any) -> Any:
    """datetime/date/time → ISO string (yyyy-mm-dd for midnight datetimes)."""
    try:
        import pandas as pd  # noqa: F401
        if isinstance(v, pd.Timestamp):
            if pd.isna(v):
                return None
            v = v.to_pydatetime()
    except Exception:
        pass
    if isinstance(v, _dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0 and v.microsecond == 0:
            return v.date().isoformat()
        return v.strftime("%Y-%m-%d %H:%M" if v.second == 0 else "%Y-%m-%d %H:%M:%S")
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, _dt.time):
        return v.strftime("%H:%M" if v.second == 0 else "%H:%M:%S")
    if isinstance(v, _dt.timedelta):
        return str(v)
    return v


def _is_datelike(v: Any) -> bool:
    try:
        import pandas as pd
        if isinstance(v, pd.Timestamp):
            return True
    except Exception:
        pass
    return isinstance(v, (_dt.date, _dt.datetime))


def fmt_num(x: Any) -> str:
    """Compact number formatting for tables and stats (no thousands separators)."""
    if not _is_number(x):
        return "" if x is None else str(x)
    if isinstance(x, int):
        return str(x)
    if math.isinf(x):
        return "inf" if x > 0 else "-inf"
    if float(x).is_integer() and abs(x) < 1e15:
        return str(int(x))
    if abs(x) >= 1e15 or (abs(x) < 1e-4 and x != 0):
        return f"{x:.6g}"
    return f"{x:.4f}".rstrip("0").rstrip(".")


def json_value(v: Any) -> Any:
    """Make a DataFrame cell JSON/markdown friendly (numpy → python, NaN → None)."""
    if v is None:
        return None
    try:
        import numpy as np
        if isinstance(v, np.generic):
            v = v.item()
    except Exception:
        pass
    if isinstance(v, float):
        if math.isnan(v):
            return None
        if math.isinf(v):
            return None
        if v.is_integer() and abs(v) < 2**53:
            return int(v)
        return round(v, 6)
    if isinstance(v, (int, str, bool)):
        return v
    v2 = iso_value(v)
    if v2 is not v:
        return v2
    return str(v)


# ── Formula fallback evaluator (cached value missing) ────────────────────────

_TOKEN_RE = re.compile(
    r"\s*(?:"
    r"(?P<num>\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)"
    r"|(?P<range>\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?\d+)"
    r"|(?P<func>[A-Za-z][A-Za-z0-9.]*)\s*\("
    r"|(?P<ref>\$?[A-Za-z]{1,3}\$?\d+)"
    r"|(?P<op>[-+*/(),])"
    r")"
)
_REF_RE = re.compile(r"\$?([A-Za-z]{1,3})\$?(\d+)")


class _FormulaUnsupported(Exception):
    pass


class _FormulaEvaluator:
    """Tiny safe evaluator: + - * / ( ), numbers, same-sheet refs/ranges and
    SUM/AVERAGE/MIN/MAX/COUNT/ROUND/ABS. No eval(); anything else → None."""

    FUNCS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "ROUND", "ABS"}

    def __init__(self, grid: list[list[Any]], formulas: dict[tuple[int, int], str]):
        self.grid = grid
        self.formulas = formulas
        self.cache: dict[tuple[int, int], Any] = {}
        self.stack: set[tuple[int, int]] = set()

    def cell(self, r: int, c: int) -> Any:
        if (r, c) in self.cache:
            return self.cache[(r, c)]
        v = self.grid[r][c] if 0 <= r < len(self.grid) and 0 <= c < len(self.grid[r]) else None
        if v is None and (r, c) in self.formulas:
            if (r, c) in self.stack or len(self.stack) > 50:
                raise _FormulaUnsupported("cycle")
            self.stack.add((r, c))
            try:
                v = self.evaluate(self.formulas[(r, c)])
            finally:
                self.stack.discard((r, c))
            self.cache[(r, c)] = v
        return v

    def evaluate(self, formula: str) -> Any:
        text = (formula or "").strip()
        if text.startswith("="):
            text = text[1:]
        if "!" in text or '"' in text:
            raise _FormulaUnsupported("cross-sheet or text formula")
        tokens: list[tuple[str, str]] = []
        pos = 0
        while pos < len(text):
            if text[pos:].strip() == "":
                break
            m = _TOKEN_RE.match(text, pos)
            if not m or m.end() == pos:
                raise _FormulaUnsupported(f"token at {pos}")
            kind = m.lastgroup or ""
            tokens.append((kind, m.group(kind)))
            pos = m.end()
        self._toks = tokens
        self._i = 0
        val = self._expr()
        if self._i != len(self._toks):
            raise _FormulaUnsupported("trailing tokens")
        return val

    # recursive descent
    def _peek(self) -> tuple[str, str]:
        return self._toks[self._i] if self._i < len(self._toks) else ("", "")

    def _take(self) -> tuple[str, str]:
        t = self._peek()
        self._i += 1
        return t

    @staticmethod
    def _num(v: Any) -> float:
        if v is None:
            return 0
        if _is_number(v):
            return v
        if isinstance(v, bool):
            return int(v)
        raise _FormulaUnsupported("non-numeric operand")

    def _expr(self) -> Any:
        v = self._term()
        while self._peek() in (("op", "+"), ("op", "-")):
            op = self._take()[1]
            rhs = self._term()
            v = self._num(v) + self._num(rhs) if op == "+" else self._num(v) - self._num(rhs)
        return v

    def _term(self) -> Any:
        v = self._factor()
        while self._peek() in (("op", "*"), ("op", "/")):
            op = self._take()[1]
            rhs = self._factor()
            if op == "*":
                v = self._num(v) * self._num(rhs)
            else:
                d = self._num(rhs)
                if d == 0:
                    raise _FormulaUnsupported("division by zero")
                v = self._num(v) / d
        return v

    def _factor(self) -> Any:
        kind, val = self._take()
        if kind == "op" and val == "-":
            return -self._num(self._factor())
        if kind == "op" and val == "+":
            return self._num(self._factor())
        if kind == "num":
            f = float(val)
            return int(f) if f.is_integer() and "." not in val and "e" not in val.lower() else f
        if kind == "ref":
            m = _REF_RE.match(val)
            return self.cell(int(m.group(2)) - 1, _letter_index(m.group(1)))
        if kind == "op" and val == "(":
            v = self._expr()
            if self._take() != ("op", ")"):
                raise _FormulaUnsupported("missing )")
            return v
        if kind == "func":
            return self._call(val.upper())
        raise _FormulaUnsupported(f"unexpected {val!r}")

    def _range_values(self, spec: str) -> list[Any]:
        a, b = spec.split(":")
        ma, mb = _REF_RE.match(a), _REF_RE.match(b)
        r1, r2 = sorted((int(ma.group(2)) - 1, int(mb.group(2)) - 1))
        c1, c2 = sorted((_letter_index(ma.group(1)), _letter_index(mb.group(1))))
        if (r2 - r1 + 1) * (c2 - c1 + 1) > 200_000:
            raise _FormulaUnsupported("range too large")
        return [self.cell(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    def _call(self, name: str) -> Any:
        if name not in self.FUNCS:
            raise _FormulaUnsupported(f"function {name}")
        args: list[list[Any]] = []
        if self._peek() != ("op", ")"):
            while True:
                kind, val = self._peek()
                if kind == "range":
                    self._take()
                    args.append(self._range_values(val))
                else:
                    args.append([self._expr()])
                if self._peek() == ("op", ","):
                    self._take()
                    continue
                break
        if self._take() != ("op", ")"):
            raise _FormulaUnsupported("missing )")
        flat = [v for a in args for v in a]
        nums = [v for v in flat if _is_number(v)]
        if name == "SUM":
            return sum(nums)
        if name == "COUNT":
            return len(nums)
        if name == "AVERAGE":
            if not nums:
                raise _FormulaUnsupported("average of nothing")
            return sum(nums) / len(nums)
        if name == "MIN":
            return min(nums) if nums else 0
        if name == "MAX":
            return max(nums) if nums else 0
        if name == "ABS":
            return abs(self._num(flat[0] if flat else 0))
        if name == "ROUND":
            x = self._num(flat[0] if flat else 0)
            n = int(self._num(flat[1])) if len(flat) > 1 else 0
            return round(x, n)
        raise _FormulaUnsupported(name)


# ── Raw loaders: every format → list of raw sheets ───────────────────────────

class RawSheet:
    __slots__ = ("name", "rows", "hidden", "formula_cols", "uncomputed")

    def __init__(self, name: str, rows: list[list[Any]], hidden: bool = False,
                 formula_cols: Optional[set[int]] = None, uncomputed: int = 0):
        self.name = name
        self.rows = rows
        self.hidden = hidden
        self.formula_cols = formula_cols or set()
        self.uncomputed = uncomputed


def _fill_merged(grid: list[list[Any]], ranges: Iterable[tuple[int, int, int, int]]) -> None:
    """Forward-fill merged ranges (0-based inclusive r1, c1, r2, c2) with the
    top-left value."""
    for r1, c1, r2, c2 in ranges:
        if r1 >= len(grid):
            continue
        top = grid[r1][c1] if c1 < len(grid[r1]) else None
        if _is_empty(top):
            continue
        for r in range(r1, min(r2, len(grid) - 1) + 1):
            row = grid[r]
            for c in range(c1, c2 + 1):
                if c >= len(row):
                    row.extend([None] * (c + 1 - len(row)))
                if _is_empty(row[c]):
                    row[c] = top


def _load_openpyxl(path: Path) -> list[RawSheet]:
    from openpyxl import load_workbook  # type: ignore[import-not-found]

    size = path.stat().st_size
    read_only = size > MERGED_CELLS_MAX_BYTES
    formulas_by_sheet: dict[str, dict[tuple[int, int], str]] = {}
    if size <= FORMULA_SCAN_MAX_BYTES:
        try:
            wbf = load_workbook(path, data_only=False, read_only=False)
            for ws in wbf.worksheets:
                found: dict[tuple[int, int], str] = {}
                for row in ws.iter_rows():
                    for cell in row:
                        if getattr(cell, "data_type", None) == "f":
                            v = cell.value
                            text = v if isinstance(v, str) else getattr(v, "text", "") or ""
                            found[(cell.row - 1, cell.column - 1)] = str(text)
                formulas_by_sheet[ws.title] = found
            wbf.close()
        except Exception as exc:  # formula detection is best-effort
            log.debug("formula scan failed for %s: %s", path.name, exc)

    wb = load_workbook(path, data_only=True, read_only=read_only)
    out: list[RawSheet] = []
    try:
        for ws in wb.worksheets:
            grid = [list(r) for r in ws.iter_rows(values_only=True)]
            formulas = formulas_by_sheet.get(ws.title, {})
            uncomputed = 0
            if formulas:
                ev = _FormulaEvaluator(grid, formulas)
                for (r, c), ftxt in formulas.items():
                    if r < len(grid) and c < len(grid[r]) and grid[r][c] is None:
                        try:
                            val = ev.cell(r, c)
                            if isinstance(val, float):
                                val = round(val, 10)
                            grid[r][c] = val
                        except Exception:
                            uncomputed += 1
            merged = None if read_only else getattr(ws, "merged_cells", None)
            if merged is not None:
                _fill_merged(grid, [
                    (m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1)
                    for m in merged.ranges
                ])
            state = getattr(ws, "sheet_state", "visible") or "visible"
            out.append(RawSheet(
                ws.title, grid, hidden=state != "visible",
                formula_cols={c for (_r, c) in formulas}, uncomputed=uncomputed,
            ))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def _load_xls(path: Path) -> list[RawSheet]:
    import xlrd  # type: ignore[import-not-found]

    try:
        book = xlrd.open_workbook(str(path), formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(str(path))
    out: list[RawSheet] = []
    for sh in book.sheets():
        grid: list[list[Any]] = []
        for r in range(sh.nrows):
            row: list[Any] = []
            for c in range(sh.ncols):
                ctype = sh.cell_type(r, c)
                v = sh.cell_value(r, c)
                if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                elif ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    except Exception:
                        pass
                elif ctype == xlrd.XL_CELL_NUMBER:
                    if isinstance(v, float) and v.is_integer() and abs(v) < 2**53:
                        v = int(v)
                elif ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif ctype == xlrd.XL_CELL_ERROR:
                    v = None
                row.append(v)
            grid.append(row)
        merged = [(rlo, clo, rhi - 1, chi - 1) for (rlo, rhi, clo, chi) in (getattr(sh, "merged_cells", None) or [])]
        _fill_merged(grid, merged)
        out.append(RawSheet(sh.name, grid, hidden=bool(getattr(sh, "visibility", 0))))
    return out


def _load_ods(path: Path) -> list[RawSheet]:
    import pandas as pd  # type: ignore[import-not-found]

    frames = pd.read_excel(path, engine="odf", sheet_name=None, header=None)
    out: list[RawSheet] = []
    for name, df in frames.items():
        grid: list[list[Any]] = []
        for row in df.itertuples(index=False, name=None):
            vals: list[Any] = []
            for v in row:
                if _is_empty(v):
                    vals.append(None)
                elif isinstance(v, pd.Timestamp):
                    vals.append(v.to_pydatetime())
                elif _is_number(v):
                    vals.append(int(v) if isinstance(v, float) and v.is_integer() and abs(v) < 2**53 else v)
                else:
                    vals.append(v.item() if hasattr(v, "item") else v)
            grid.append(vals)
        out.append(RawSheet(str(name), grid))
    return out


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _load_delimited(path: Path, sep: Optional[str]) -> list[RawSheet]:
    text = _read_text(path)
    if sep is None:
        sample = text[:20000]
        try:
            sep = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            sep = ","
    grid: list[list[Any]] = []
    reader = csv.reader(text.splitlines(), delimiter=sep)
    for row in reader:
        grid.append([None if not (c or "").strip() else coerce_text_number(c) for c in row])
    return [RawSheet(path.stem or "Sheet1", grid)]


def load_raw_sheets(path: Path) -> list[RawSheet]:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _load_openpyxl(path)
    if ext == ".xls":
        return _load_xls(path)
    if ext == ".ods":
        return _load_ods(path)
    if ext == ".tsv":
        return _load_delimited(path, "\t")
    if ext == ".csv":
        return _load_delimited(path, None)
    raise ValueError(f"Unsupported spreadsheet type: {ext or path.name}")


# ── Raw grid → DataFrame ─────────────────────────────────────────────────────

def _looks_numeric_text(v: Any) -> bool:
    return isinstance(v, str) and bool(_NUM_RE.match(v.strip()) or _THOUSANDS_RE.match(v.strip()))


def _is_header_row(row: list[Any], rows_after: int = 1) -> bool:
    """A header candidate: text cells cover ≥ 50% of the columns and hold ≥ 2
    distinct values. A 1-column sheet has no second value to compare, so its
    lone text cell only counts as a header when ≥ 1 data row follows it — a
    single-cell sheet (or a lone value) is data, never a header."""
    width = len(row)
    if width == 0:
        return False
    texts = [str(v).strip() for v in row if isinstance(v, str) and v.strip() and not _looks_numeric_text(v)]
    if len(texts) * 2 < width:
        return False
    if width == 1:
        return rows_after >= 1
    return len(set(texts)) >= 2


def _dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        key = n.lower()
        if key in seen:
            seen[key] += 1
            cand = f"{n} ({seen[key]})"
            while cand.lower() in seen:
                seen[key] += 1
                cand = f"{n} ({seen[key]})"
            seen[cand.lower()] = 1
            out.append(cand)
        else:
            seen[key] = 1
            out.append(n)
    return out


def _column_values(values: list[Any]) -> tuple[list[Any], str]:
    """Normalise one column's values and give it a dtype label:
    number | date | bool | text | mixed | empty."""
    non_null = [v for v in values if not _is_empty(v)]
    if not non_null:
        return [None] * len(values), "empty"
    if all(isinstance(v, bool) for v in non_null):
        return [None if _is_empty(v) else v for v in values], "bool"
    if all(_is_number(v) for v in non_null):
        return [None if _is_empty(v) else v for v in values], "number"
    if all(_is_datelike(v) for v in non_null):
        return [None if _is_empty(v) else iso_value(v) for v in values], "date"
    if all(isinstance(v, str) for v in non_null):
        cleaned = [None if _is_empty(v) else v.strip() for v in values]
        if all(_ISO_DATE_RE.match(v) for v in cleaned if v is not None):
            return cleaned, "date"
        return cleaned, "text"
    out: list[Any] = []
    for v in values:
        if _is_empty(v):
            out.append(None)
        elif _is_number(v):
            out.append(v)
        elif isinstance(v, str):
            out.append(v.strip())
        else:
            iv = iso_value(v)
            out.append(iv if isinstance(iv, str) else str(v))
    return out, "mixed"


def build_sheet(raw: RawSheet) -> dict[str, Any]:
    """Raw grid → {"name", "hidden", "df", "dtypes", "formula_columns",
    "header_row", "preamble", "uncomputed_formulas"}."""
    import pandas as pd  # type: ignore[import-not-found]

    grid = [list(r) for r in (raw.rows or [])]
    width = max((len(r) for r in grid), default=0)
    for r in grid:
        if len(r) < width:
            r.extend([None] * (width - len(r)))
    # Keep non-empty rows/columns (with their original sheet positions)
    keep_rows = [i for i, r in enumerate(grid) if any(not _is_empty(v) for v in r)]
    keep_cols = [c for c in range(width) if any(not _is_empty(grid[i][c]) for i in keep_rows)]
    rows = [[grid[i][c] for c in keep_cols] for i in keep_rows]

    header_idx: Optional[int] = None
    for i, r in enumerate(rows[:HEADER_SCAN_ROWS]):
        if _is_header_row(r, len(rows) - i - 1):
            header_idx = i
            break

    preamble: list[str] = []
    if header_idx is not None:
        for r in rows[:header_idx]:
            vals = []
            for v in r:
                s = "" if _is_empty(v) else str(iso_value(v))
                if s and s not in vals:
                    vals.append(s)
            if vals:
                preamble.append(" · ".join(vals)[:200])
        head = rows[header_idx]
        names = []
        for j, v in enumerate(head):
            s = "" if _is_empty(v) else str(iso_value(v)).strip()
            s = re.sub(r"\s+", " ", s)
            names.append(s or f"Column {col_letter(keep_cols[j])}")
        data = rows[header_idx + 1:]
        header_row = keep_rows[header_idx] + 1
    else:
        names = [f"Column {col_letter(c)}" for c in keep_cols]
        data = rows
        header_row = None
    names = _dedupe(names)

    columns: dict[str, list[Any]] = {}
    dtypes: dict[str, str] = {}
    for j, name in enumerate(names):
        vals, dtype = _column_values([r[j] for r in data])
        columns[name] = vals
        dtypes[name] = dtype

    df = pd.DataFrame(columns, columns=names)
    for name, dtype in dtypes.items():
        if dtype == "number":
            df[name] = pd.to_numeric(df[name], errors="coerce")
        elif dtype == "empty":
            df[name] = df[name].astype(object)
    # Drop rows that became empty under the header (e.g. data all blank)
    if len(df):
        df = df.dropna(how="all").reset_index(drop=True)

    formula_columns = [names[j] for j, c in enumerate(keep_cols) if c in raw.formula_cols]
    return {
        "name": raw.name,
        "hidden": bool(raw.hidden),
        "df": df,
        "dtypes": dtypes,
        "formula_columns": formula_columns,
        "header_row": header_row,
        "preamble": preamble[:5],
        "uncomputed_formulas": int(raw.uncomputed),
    }


def parse_workbook(path: Path) -> list[dict[str, Any]]:
    return [build_sheet(raw) for raw in load_raw_sheets(Path(path))]


# ── Text rendering ───────────────────────────────────────────────────────────

def _md_cell(v: Any) -> str:
    v = json_value(v)
    if v is None:
        return ""
    s = fmt_num(v) if _is_number(v) else str(v)
    return s.replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()


def markdown_table(columns: list[str], rows: Iterable[Iterable[Any]],
                   numeric: Optional[set[str]] = None) -> str:
    numeric = numeric or set()
    head = "| " + " | ".join(_md_cell(c) or " " for c in columns) + " |"
    sep = "| " + " | ".join("---:" if c in numeric else "---" for c in columns) + " |"
    body = ["| " + " | ".join(_md_cell(v) for v in r) + " |" for r in rows]
    return "\n".join([head, sep] + body)


def numeric_columns(sheet: dict[str, Any]) -> list[str]:
    return [c for c, t in (sheet.get("dtypes") or {}).items() if t == "number" and c in sheet["df"].columns]


def column_stats(sheet: dict[str, Any]) -> dict[str, dict[str, Any]]:
    df = sheet["df"]
    out: dict[str, dict[str, Any]] = {}
    for c in numeric_columns(sheet):
        s = df[c].dropna()
        if s.empty:
            continue
        out[c] = {
            "sum": json_value(s.sum()),
            "mean": json_value(s.mean()),
            "min": json_value(s.min()),
            "max": json_value(s.max()),
        }
    return out


def stats_line(sheet: dict[str, Any], max_cols: int = 30) -> str:
    stats = column_stats(sheet)
    if not stats:
        return ""
    parts = []
    for c, st in list(stats.items())[:max_cols]:
        parts.append(
            f"{c} — sum {fmt_num(st['sum'])}, mean {fmt_num(st['mean'])}, "
            f"min {fmt_num(st['min'])}, max {fmt_num(st['max'])}"
        )
    more = len(stats) - max_cols
    return "Stats (all rows): " + "; ".join(parts) + (f"; … {more} more numeric columns" if more > 0 else "")


def sheet_title(sheet: dict[str, Any]) -> str:
    df = sheet["df"]
    return f"### Sheet: {sheet['name']} ({len(df)} rows × {len(df.columns)} columns)"


def render_sheet_page(sheet: dict[str, Any], row_cap: int = TEXT_ROW_CAP) -> str:
    df = sheet["df"]
    lines = [sheet_title(sheet)]
    for p in sheet.get("preamble") or []:
        lines.append(p)
    if len(df.columns) == 0:
        lines.append("(empty sheet)")
        return "\n".join(lines)
    num = set(numeric_columns(sheet))
    lines.append("")
    lines.append(markdown_table(list(df.columns), df.head(row_cap).itertuples(index=False, name=None), num))
    if len(df) > row_cap:
        lines.append(f"… {len(df) - row_cap} more rows (showing the first {row_cap} of {len(df)})")
    st = stats_line(sheet)
    if st:
        lines.append("")
        lines.append(st)
    if sheet.get("formula_columns"):
        lines.append("Formula columns: " + ", ".join(sheet["formula_columns"]))
    return "\n".join(lines)


def sheet_meta(sheet: dict[str, Any]) -> dict[str, Any]:
    df = sheet["df"]
    if sheet.get("hidden"):
        # never expose a hidden sheet's columns / samples
        return {"name": sheet["name"], "hidden": True, "rows": int(len(df)), "cols": int(len(df.columns))}
    cols = []
    formula = set(sheet.get("formula_columns") or [])
    for c in list(df.columns)[:META_MAX_COLUMNS]:
        s = df[c]
        sample = [json_value(v) for v in s.dropna().head(3).tolist()]
        entry = {
            "name": c,
            "dtype": (sheet.get("dtypes") or {}).get(c, "text"),
            "non_null": int(s.notna().sum()),
            "sample": sample,
        }
        if c in formula:
            entry["formula"] = True
        cols.append(entry)
    return {
        "name": sheet["name"],
        "rows": int(len(df)),
        "cols": int(len(df.columns)),
        "columns": cols,
        "hidden": bool(sheet.get("hidden")),
        "header_row": sheet.get("header_row"),
    }


def build_meta(sheets: list[dict[str, Any]], fmt: str) -> dict[str, Any]:
    visible = [s for s in sheets if not s.get("hidden")]
    meta: dict[str, Any] = {
        "sheets": [sheet_meta(s) for s in sheets],
        "sheet_count": len(visible),
        "hidden_sheets": [s["name"] for s in sheets if s.get("hidden")],
        "table_rows": int(sum(len(s["df"]) for s in visible)),
        "sheet_format": fmt,
    }
    unc = sum(int(s.get("uncomputed_formulas") or 0) for s in sheets)
    if unc:
        meta["uncomputed_formulas"] = unc
    return meta


def normalize_meta_sheets(meta: Any) -> tuple[list[dict[str, Any]], bool]:
    """meta["sheets"] as a list, whatever its stored form: the current list,
    or the legacy {name: {"rows", "cols"}} map of older uploads. Hidden
    entries are reduced to {name, hidden, rows, cols}. Returns (list,
    changed)."""
    raw = meta.get("sheets") if isinstance(meta, dict) else None
    changed = False
    out: list[dict[str, Any]] = []
    if isinstance(raw, dict):
        changed = True
        for name, info in raw.items():
            info = info if isinstance(info, dict) else {}
            out.append({
                "name": str(name),
                "rows": int(info.get("rows") or 0),
                "cols": int(info.get("cols") or 0),
                "hidden": bool(info.get("hidden")),
            })
    elif isinstance(raw, list):
        for entry in raw:
            if not isinstance(entry, dict):
                changed = True
                continue
            if entry.get("hidden") and set(entry) - {"name", "hidden", "rows", "cols"}:
                changed = True
                entry = {"name": str(entry.get("name", "")), "hidden": True,
                         "rows": int(entry.get("rows") or 0), "cols": int(entry.get("cols") or 0)}
            out.append(entry)
    elif raw is not None:
        changed = True
    return out, changed


def normalize_meta(meta: Any) -> bool:
    """Normalise meta["sheets"] in place (see normalize_meta_sheets); fills
    sheet_count / hidden_sheets / table_rows when missing. True if changed."""
    if not isinstance(meta, dict) or "sheets" not in meta:
        return False
    sheets, changed = normalize_meta_sheets(meta)
    if not changed:
        return False
    meta["sheets"] = sheets
    visible = [x for x in sheets if not x.get("hidden")]
    meta["sheet_count"] = len(visible)
    meta["hidden_sheets"] = [x.get("name") for x in sheets if x.get("hidden")]
    meta["table_rows"] = int(sum(int(x.get("rows") or 0) for x in visible))
    return True


def render_pages(sheets: list[dict[str, Any]], row_cap: int = TEXT_ROW_CAP) -> list[str]:
    pages = [render_sheet_page(s, row_cap) for s in sheets if not s.get("hidden")]
    return pages or ["(workbook has no visible sheets)"]


# ── Cache ────────────────────────────────────────────────────────────────────

_MEM_CACHE: dict[str, tuple[float, list[dict[str, Any]]]] = {}
_MEM_LOCK = threading.Lock()
_MEM_MAX = 8


def _write_cache(cache_path: Path, sheets: list[dict[str, Any]]) -> None:
    try:
        tmp = cache_path.with_suffix(".tmp")
        with open(tmp, "wb") as fh:
            pickle.dump({"version": CACHE_VERSION, "sheets": sheets}, fh, protocol=pickle.HIGHEST_PROTOCOL)
        tmp.replace(cache_path)
    except Exception as exc:
        log.debug("sheet cache write failed (%s): %s", cache_path, exc)


def _read_cache(cache_path: Path) -> Optional[list[dict[str, Any]]]:
    try:
        with open(cache_path, "rb") as fh:
            payload = pickle.load(fh)
        if isinstance(payload, dict) and payload.get("version") == CACHE_VERSION:
            sheets = payload.get("sheets")
            if isinstance(sheets, list):
                return sheets
    except Exception as exc:
        log.debug("sheet cache read failed (%s): %s", cache_path, exc)
    return None


def _remember(key: str, mtime: float, sheets: list[dict[str, Any]]) -> None:
    with _MEM_LOCK:
        _MEM_CACHE[key] = (mtime, sheets)
        while len(_MEM_CACHE) > _MEM_MAX:
            _MEM_CACHE.pop(next(iter(_MEM_CACHE)))


def extract(path: Path, cache_path: Optional[Path] = None) -> tuple[list[str], dict[str, Any]]:
    """Parse a spreadsheet → (page_texts, meta); writes the frame cache when
    `cache_path` is given."""
    path = Path(path)
    sheets = parse_workbook(path)
    if cache_path is not None:
        _write_cache(Path(cache_path), sheets)
        try:
            _remember(str(cache_path), Path(cache_path).stat().st_mtime, sheets)
        except OSError:
            pass
    fmt = path.suffix.lower().lstrip(".") or "sheet"
    return render_pages(sheets), build_meta(sheets, fmt)


def load_sheets(raw_path: Path, cache_path: Optional[Path] = None) -> list[dict[str, Any]]:
    """Parsed sheets for a raw file: memory cache → pickle cache → re-parse
    (and rewrite the pickle). Raises when the raw file can't be parsed."""
    raw_path = Path(raw_path)
    if cache_path is None:
        cache_path = raw_path.parent / CACHE_NAME
    cache_path = Path(cache_path)
    raw_mtime = raw_path.stat().st_mtime if raw_path.exists() else 0.0
    try:
        cache_mtime = cache_path.stat().st_mtime
    except OSError:
        cache_mtime = None
    key = str(cache_path)
    if cache_mtime is not None and cache_mtime >= raw_mtime:
        with _MEM_LOCK:
            hit = _MEM_CACHE.get(key)
        if hit and hit[0] == cache_mtime:
            return hit[1]
        sheets = _read_cache(cache_path)
        if sheets is not None:
            _remember(key, cache_mtime, sheets)
            return sheets
    if not raw_path.exists():
        raise FileNotFoundError(str(raw_path))
    sheets = parse_workbook(raw_path)
    _write_cache(cache_path, sheets)
    try:
        _remember(key, cache_path.stat().st_mtime, sheets)
    except OSError:
        pass
    return sheets


def forget(cache_path: Path) -> None:
    with _MEM_LOCK:
        _MEM_CACHE.pop(str(cache_path), None)


# ── Preview (GET /api/idp/documents/{id}/sheets) ─────────────────────────────

PREVIEW_MAX_LIMIT = 1000


def preview(sheets: list[dict[str, Any]], sheet: Optional[str] = None,
            offset: int = 0, limit: int = 200, include_hidden: bool = False) -> dict[str, Any]:
    """One page of rows of one sheet (default: the first visible sheet).
    Hidden sheets are neither listed nor served unless `include_hidden`;
    asking for one by name raises PermissionError("sheet is hidden")."""
    listed = [s for s in sheets if include_hidden or not s.get("hidden")]
    hidden_count = sum(1 for s in sheets if s.get("hidden"))
    listing = [
        {"name": s["name"], "rows": int(len(s["df"])), "cols": int(len(s["df"].columns)),
         "hidden": bool(s.get("hidden"))}
        for s in listed
    ]
    empty = {"sheets": listing, "sheet": None, "columns": [], "dtypes": [], "rows": [],
             "offset": 0, "total": 0, "formula_columns": [], "hidden_sheets": hidden_count}
    if not sheets:
        return empty
    chosen = None
    if sheet:
        def _find(pool: list[dict[str, Any]]) -> Optional[dict[str, Any]]:
            return next((s for s in pool if s["name"] == sheet), None) or next(
                (s for s in pool if s["name"].lower() == sheet.strip().lower()), None)
        chosen = _find(listed)
        if chosen is None:
            if not include_hidden and _find([s for s in sheets if s.get("hidden")]) is not None:
                raise PermissionError("sheet is hidden")
            raise KeyError(sheet)
    if chosen is None:
        chosen = next((s for s in sheets if not s.get("hidden")), None)
        if chosen is None:
            if not include_hidden:
                return empty
            chosen = sheets[0]
    df = chosen["df"]
    offset = max(0, int(offset or 0))
    limit = max(1, min(int(limit or 200), PREVIEW_MAX_LIMIT))
    part = df.iloc[offset: offset + limit]
    rows = [[json_value(v) for v in r] for r in part.itertuples(index=False, name=None)]
    dtypes = chosen.get("dtypes") or {}
    return {
        "sheets": listing,
        "sheet": chosen["name"],
        "columns": [str(c) for c in df.columns],
        "dtypes": [dtypes.get(c, "text") for c in df.columns],
        "rows": rows,
        "offset": offset,
        "total": int(len(df)),
        "formula_columns": list(chosen.get("formula_columns") or []),
        "hidden_sheets": hidden_count,
    }
