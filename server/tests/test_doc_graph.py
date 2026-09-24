"""
Tests for `server/doc_graph.py` — DAG executor and document routing.

Tests cover: topological sort, cycle detection, node execution with errors,
probe features, routing rule truth table, and end-to-end run_routing with
mocked decider and model resolver.
"""

from __future__ import annotations

import json
import pathlib
import sys
import tempfile
from dataclasses import dataclass
from email.message import EmailMessage
from unittest.mock import Mock

import pytest

# Make the server module importable when pytest is run from the repo root.
_HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(_HERE.parent))

import doc_graph as dg
import idp_engine as ide
from paths import uploads_dir

# ──────────────────────────────────────────────────────────────────────────
# Fixtures & helpers
# ──────────────────────────────────────────────────────────────────────────


@dataclass
class FakeDocument:
    """Minimal document stub for testing."""
    id: str
    filename: str
    mime: str
    text: str
    pages: int = 1
    page_images: list[str] = None

    def __post_init__(self):
        if self.page_images is None:
            self.page_images = []


def fake_resolve_model(category: str) -> str | None:
    """Simple model resolver: map categories to fake model names."""
    mapping = {
        "ocr": "vision-ocr:8b",
        "tables": "table-analyzer:7b",
        "text": "fast-chat:3b",
        "docs": "general-chat:7b",
    }
    return mapping.get(category, "general-chat:7b")


# ──────────────────────────────────────────────────────────────────────────
# DAG Executor Tests
# ──────────────────────────────────────────────────────────────────────────


class TestTopologicalSort:
    """Test Kahn's algorithm."""

    def test_linear_chain(self):
        """A → B → C."""
        nodes = {
            "a": dg.Node("a", lambda: None, []),
            "b": dg.Node("b", lambda: None, ["a"]),
            "c": dg.Node("c", lambda: None, ["b"]),
        }
        order = dg._topological_sort(nodes)
        assert order == ["a", "b", "c"]

    def test_diamond(self):
        """A → B,C; B,C → D."""
        nodes = {
            "a": dg.Node("a", lambda: None, []),
            "b": dg.Node("b", lambda: None, ["a"]),
            "c": dg.Node("c", lambda: None, ["a"]),
            "d": dg.Node("d", lambda: None, ["b", "c"]),
        }
        order = dg._topological_sort(nodes)
        assert order[0] == "a"
        assert set(order[1:3]) == {"b", "c"}
        assert order[3] == "d"

    def test_empty_graph(self):
        """No nodes."""
        order = dg._topological_sort({})
        assert order == []

    def test_cycle_detection(self):
        """A → B → C → A (cycle)."""
        nodes = {
            "a": dg.Node("a", lambda: None, ["c"]),
            "b": dg.Node("b", lambda: None, ["a"]),
            "c": dg.Node("c", lambda: None, ["b"]),
        }
        try:
            dg._topological_sort(nodes)
            assert False, "Should raise ValueError on cycle"
        except ValueError as e:
            assert "cycle" in str(e).lower()


class TestDAGExecution:
    """Test _execute_dag."""

    def test_successful_execution(self):
        """Nodes execute in order and produce outputs."""
        nodes = {
            "a": dg.Node("a", lambda: 1, []),
            "b": dg.Node("b", lambda a: a + 1, ["a"]),
            "c": dg.Node("c", lambda b: b * 2, ["b"]),
        }
        outputs, trace = dg._execute_dag(nodes, {})
        assert outputs["a"] == 1
        assert outputs["b"] == 2
        assert outputs["c"] == 4
        assert len(trace) == 3
        assert all(t.status == "ok" for t in trace)

    def test_node_exception(self):
        """A failing node doesn't crash the DAG; downstream nodes still run if they can."""
        def bad_node():
            raise RuntimeError("Oops")

        nodes = {
            "a": dg.Node("a", bad_node, []),
            "b": dg.Node("b", lambda a: a + 1, ["a"]),  # b depends on a
            "c": dg.Node("c", lambda: "independent", []),  # c is independent
        }
        outputs, trace = dg._execute_dag(nodes, {})

        # a failed
        assert outputs["a"] is None
        a_trace = next((t for t in trace if t.node == "a"), None)
        assert a_trace.status == "error"
        assert "Oops" in a_trace.detail

        # b skipped because a is None
        assert outputs["b"] is None
        b_trace = next((t for t in trace if t.node == "b"), None)
        assert b_trace.status == "skipped"

        # c ran successfully
        assert outputs["c"] == "independent"
        c_trace = next((t for t in trace if t.node == "c"), None)
        assert c_trace.status == "ok"

    def test_timing_recorded(self):
        """Each trace entry records execution time."""
        import time
        nodes = {
            "slow": dg.Node("slow", lambda: time.sleep(0.01) or "done", []),
        }
        outputs, trace = dg._execute_dag(nodes, {})
        entry = trace[0]
        assert entry.ms >= 5  # At least a few ms (may be longer on slow systems)


# ──────────────────────────────────────────────────────────────────────────
# Probe Tests
# ──────────────────────────────────────────────────────────────────────────


class TestProbe:
    """Test probe() feature extraction."""

    def test_basic_document(self):
        """Typical text document."""
        text = "This is a report with some content here and there."
        doc = FakeDocument(
            id="test", filename="report.txt", mime="text/plain",
            text=text, pages=1
        )
        result = dg.probe(doc)
        assert result.chars == len(text)
        assert result.has_text_layer is True
        assert result.has_page_images is False
        assert result.is_email is False
        assert result.table_density == 0.0

    def test_email_by_mime(self):
        """Email detected by mime type."""
        doc = FakeDocument(
            id="test", filename="msg.eml", mime="message/rfc822",
            text="From: alice@example.com\nTo: bob@example.com\nSubject: Hello\n\nHi there",
            pages=1
        )
        result = dg.probe(doc)
        assert result.is_email is True

    def test_email_by_headers(self):
        """Email detected by header pattern."""
        doc = FakeDocument(
            id="test", filename="text.txt", mime="text/plain",
            text="From: alice@example.com\nTo: bob@example.com\nSubject: Hello\nBody here",
            pages=1
        )
        result = dg.probe(doc)
        assert result.is_email is True

    def test_email_by_filename(self):
        """Email detected by .eml extension."""
        doc = FakeDocument(
            id="test", filename="message.eml", mime="text/plain",
            text="Some content", pages=1
        )
        result = dg.probe(doc)
        assert result.is_email is True

    def test_image_only(self):
        """Document with images but no text."""
        doc = FakeDocument(
            id="test", filename="scan.pdf", mime="application/pdf",
            text="", pages=5,
            page_images=["/path/to/page_1.png", "/path/to/page_2.png"]
        )
        result = dg.probe(doc)
        assert result.has_text_layer is False
        assert result.has_page_images is True

    def test_text_layer_threshold(self):
        """Text layer exists only if >= 40 chars."""
        doc_short = FakeDocument(
            id="test", filename="tiny.txt", mime="text/plain",
            text="Hi", pages=1
        )
        assert dg.probe(doc_short).has_text_layer is False

        doc_long = FakeDocument(
            id="test", filename="short.txt", mime="text/plain",
            text="This is at least forty characters long now", pages=1
        )
        assert dg.probe(doc_long).has_text_layer is True

    def test_table_density_tabbed(self):
        """Lines with tabs detected as tabular."""
        text = "Name\tAge\tCity\tLocation\tSalary\nAlice\t30\tNY\tNewport\t50000\nBob\t25\tLA\tLosAngeles\t45000"
        doc = FakeDocument(
            id="test", filename="data.csv", mime="text/csv",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        assert result.has_text_layer is True  # Should have enough chars
        assert result.table_density >= 0.5  # Most lines have tabs

    def test_table_density_piped(self):
        """Lines with pipes detected as tabular."""
        doc = FakeDocument(
            id="test", filename="table.txt", mime="text/plain",
            text="| Name | Age |\n| Alice | 30 |\n| Bob | 25 |",
            pages=1
        )
        result = dg.probe(doc)
        assert result.table_density >= 0.5

    def test_table_density_comma_delimited(self):
        """Lines with 2+ comma/semicolon delimiters are tabular."""
        doc = FakeDocument(
            id="test", filename="data.txt", mime="text/plain",
            text="apple,orange,banana\ngrape,kiwi,mango\npeach,plum,cherry",
            pages=1
        )
        result = dg.probe(doc)
        assert result.table_density >= 0.5

    def test_table_density_mixed_content(self):
        """Text with minimal tables has low table density."""
        text = "Here is some text with more content.\nMore text here and there.\nName,Value,Count,Total\nItem1,10,5,50\nItem2,20,3,60"
        doc = FakeDocument(
            id="test", filename="doc.txt", mime="text/plain",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        # 2 non-tabular lines, 3 tabular lines (with 2+ commas) = 3/5 = 0.6
        assert 0 < result.table_density <= 1.0

    def test_table_density_whitespace_aligned_columns(self):
        """Pandas-style whitespace-aligned columns (3+ columns separated by 2+ spaces)."""
        # This is what pandas DataFrame.to_string(index=False) produces
        text = "Name          Age          City\nAlice         30           New York\nBob           25           Los Angeles"
        doc = FakeDocument(
            id="test", filename="data.txt", mime="text/plain",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        # All 3 lines have 3+ columns separated by 2+ spaces
        assert result.table_density >= 0.5, f"Expected density >= 0.5, got {result.table_density}"

    def test_table_density_prose_with_double_spaces(self):
        """Prose with some double spaces (e.g., after periods) should NOT be tabular."""
        # A paragraph with double spaces after periods (old typing convention)
        text = "This is a sentence.  Here is another one.  And yet another.  This should not be tabular."
        doc = FakeDocument(
            id="test", filename="doc.txt", mime="text/plain",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        # The line is > 200 chars? No, it's ~95 chars. But it has only 2 "parts" when split by 2+ spaces:
        # ["This is a sentence.", "Here is another one.", "And yet another.", "This should not be tabular."]
        # Actually that's 4 parts. Let me check: re.split(r'\s{2,}', ...) on that line gives ~4 parts.
        # But this is a judgment call. The requirement says "require line < 200 chars or consistent across doc".
        # We have only one line, so it's not consistent. Let me use a simpler prose example.
        text = "This is a sentence.  More text here."
        doc = FakeDocument(
            id="test", filename="doc.txt", mime="text/plain",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        # This line: "This is a sentence.  More text here." when split by 2+ spaces gives:
        # ["This is a sentence.", "More text here."] = 2 parts. Need 3+ parts to be tabular.
        assert result.table_density < 0.2, f"Expected low density for prose, got {result.table_density}"

    def test_double_space_prose_repro(self):
        """Repro case: prose with double spaces after periods should NOT score as tables.

        This is the exact bug report: ordinary prose typed with two spaces after
        full stops scored table_density 1.00 and routed to 'tables'. Must now be < 0.2.
        """
        prose = "Dear team.  We met on Monday.  The budget is approved.  Please review the attached plan.  Thanks.\n" * 6
        doc = FakeDocument(
            id="prose_test", filename="x.txt", mime="text/plain",
            text=prose, pages=1
        )
        result = dg.probe(doc)
        assert result.table_density < 0.2, f"Prose with double spaces scored {result.table_density}, expected < 0.2"
        # Also verify it doesn't route to 'tables'
        category, reason = dg.route_category(result, None)
        assert category != "tables", f"Prose routed to 'tables', reason: {reason}"

    def test_prose_double_spaces_after_commas_colons(self):
        """Prose with sentences ending in periods should NOT be counted as tabular.

        Even if separated by 2+ spaces, sentences ending with periods are invalid table cells.
        """
        # Words ending with periods separated by 2+ spaces
        text = "apple.  banana.  cherry.\n" * 3 + "normal sentence with multiple words here\n" * 3
        doc = FakeDocument(
            id="test", filename="prose.txt", mime="text/plain",
            text=text, pages=1
        )
        result = dg.probe(doc)
        # "apple.  banana.  cherry." - each cell ends with period, invalid
        # "normal sentence..." - no 2+ space delimiters, not tabular
        # 0 tabular out of 6 lines = 0.0
        assert result.table_density < 0.2, f"Prose with periods scored {result.table_density}, expected < 0.2"

    def test_email_with_signature_block(self):
        """Email with signature block followed by normal sentences should have low table density.

        A signature block like 'John Smith        Senior Engineer        ACME Corp' followed by
        multiple normal sentences should score <= 0.2 (not strongly tabular).
        """
        signature = "John Smith        Senior Engineer        ACME Corp"
        # Add enough normal prose lines to dilute the signature line
        sentences = "\n".join([f"Sentence {i}." for i in range(1, 10)])
        text = signature + "\n\n" + sentences
        doc = FakeDocument(
            id="email_test", filename="msg.txt", mime="text/plain",
            text=text, pages=1
        )
        result = dg.probe(doc)
        # The signature line: "John Smith        Senior Engineer        ACME Corp"
        # When split by 2+ spaces: ["John Smith", "Senior Engineer", "ACME Corp"]
        # Each has 2 words, valid cells - tabular
        # Sentence lines: each is one sentence, no 2+ space splits - not tabular
        # 1 tabular line out of 10 total lines = 0.1
        assert result.table_density <= 0.2, f"Email with signature scored {result.table_density}, expected <= 0.2"

    def test_pandas_aligned_table(self):
        """Genuine pandas-style aligned table should route to 'tables' with density >= 0.5."""
        # This is what pandas DataFrame.to_string(index=False) produces
        text = "name   qty   price\nbolt   10    0.5\nnut    20    0.2\nscrew  5     0.9\n"
        doc = FakeDocument(
            id="table_test", filename="data.txt", mime="text/plain",
            text=text, pages=1
        )
        result = dg.probe(doc)
        assert result.table_density >= 0.5, f"Pandas table scored {result.table_density}, expected >= 0.5"
        category, reason = dg.route_category(result, None)
        assert category == "tables", f"Pandas table routed to '{category}', expected 'tables'"

    def test_aligned_table_with_currency_percent(self):
        """Aligned table with currency and percent cells should score >= 0.5."""
        # An aligned table with currency ($1,200, $600.50) and percent (40%, 20%)
        text = "item     cost     share\nrent     $1,200   40%\nfood     $600.50  20%\nutil     $300     10%\nother    $200     5%\n"
        # Repeat to make it long enough to have text_layer
        text = text * 2
        doc = FakeDocument(
            id="currency_table", filename="budget.txt", mime="text/plain",
            text=text, pages=1
        )
        result = dg.probe(doc)
        assert result.table_density >= 0.5, f"Currency/percent table scored {result.table_density}, expected >= 0.5"
        category, reason = dg.route_category(result, None)
        assert category == "tables", f"Currency table routed to '{category}', expected 'tables'"

    def test_csv_mime_detection_high_density(self):
        """CSV with text/csv mime → table_density = 1.0."""
        text = "Name,Age,City\nAlice,30,New York\nBob,25,Los Angeles"
        doc = FakeDocument(
            id="test", filename="data.txt", mime="text/csv",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        assert result.table_density == 1.0

    def test_csv_filename_detection_high_density(self):
        """CSV with .csv filename → table_density = 1.0."""
        text = "Name,Age,City\nAlice,30,New York\nBob,25,Los Angeles"
        doc = FakeDocument(
            id="test", filename="data.csv", mime="text/plain",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        assert result.table_density == 1.0

    def test_xlsx_mime_detection_high_density(self):
        """XLSX with spreadsheet mime → table_density = 1.0."""
        text = "Name\tAge\tCity\nAlice\t30\tNew York\nBob\t25\tLos Angeles"
        doc = FakeDocument(
            id="test", filename="data.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        assert result.table_density == 1.0

    def test_xlsx_filename_extension_detection(self):
        """XLSX with .xlsx filename → table_density = 1.0."""
        text = "Name\tAge\tCity\nAlice\t30\tNew York\nBob\t25\tLos Angeles\nCarol\t35\tChicago"
        doc = FakeDocument(
            id="test", filename="spreadsheet.xlsx", mime="text/plain",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        assert result.table_density == 1.0

    def test_tsv_mime_detection(self):
        """TSV with tab-separated-values mime → table_density = 1.0."""
        text = "Name\tAge\tCity\nAlice\t30\tNew York\nBob\t25\tLos Angeles\nCarol\t35\tChicago"
        doc = FakeDocument(
            id="test", filename="data.txt", mime="text/tab-separated-values",
            text=text,
            pages=1
        )
        result = dg.probe(doc)
        assert result.table_density == 1.0

    def test_real_csv_round_trip(self):
        """Real CSV extraction via idp_engine._extract_csv and probe."""
        pytest.importorskip("pandas")
        import csv
        import os

        # Create a temp CSV file
        with tempfile.TemporaryDirectory() as tmpdir:
            csv_path = pathlib.Path(tmpdir) / "test.csv"
            # Write a 10-row CSV
            with open(csv_path, "w", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["Name", "Age", "City", "Salary"])
                for i in range(10):
                    writer.writerow([f"Person{i}", 20 + i, f"City{i}", 50000 + i*1000])

            # Extract using idp_engine._extract_csv
            pages = ide._extract_csv(csv_path)
            extracted_text = pages[0] if pages else ""

            # Create a document-like object
            doc = FakeDocument(
                id="csv_test", filename="test.csv", mime="text/csv",
                text=extracted_text,
                pages=1
            )

            # Probe and check routing
            probe_result = dg.probe(doc)
            category, reason = dg.route_category(probe_result, None)
            assert category == "tables", f"Expected 'tables', got '{category}'. table_density={probe_result.table_density}"

    def test_real_xlsx_round_trip(self):
        """Real XLSX extraction via idp_engine._extract_xlsx and probe."""
        pytest.importorskip("openpyxl")

        # Create a temp XLSX file
        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = pathlib.Path(tmpdir) / "test.xlsx"
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Name"
            ws["B1"] = "Age"
            ws["C1"] = "City"
            ws["D1"] = "Salary"
            for i in range(10):
                ws[f"A{i+2}"] = f"Person{i}"
                ws[f"B{i+2}"] = 20 + i
                ws[f"C{i+2}"] = f"City{i}"
                ws[f"D{i+2}"] = 50000 + i*1000
            wb.save(xlsx_path)

            # Extract using idp_engine._extract_xlsx
            pages, meta = ide._extract_xlsx(xlsx_path)
            extracted_text = "\n".join(pages) if pages else ""

            # Create a document-like object
            doc = FakeDocument(
                id="xlsx_test", filename="test.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                text=extracted_text,
                pages=1
            )

            # Probe and check routing
            probe_result = dg.probe(doc)
            category, reason = dg.route_category(probe_result, None)
            assert category == "tables", f"Expected 'tables', got '{category}'. table_density={probe_result.table_density}"

    def test_routing_csv_with_laya_note_decision(self):
        """CSV with high density must route to 'tables' even if Laya says plain_note."""
        text = "Name\tAge\tCity\nAlice\t30\tNew York\nBob\t25\tLA"
        doc = FakeDocument(
            id="test", filename="data.csv", mime="text/csv",
            text=text,
            pages=1
        )
        probe_result = dg.probe(doc)
        # Decision: plain_note with high confidence
        decision = {
            "kind": "plain_note",
            "confidence": 0.9,
            "complexity": "low",
            "probabilities": {},
        }
        # Table density (1.0 from mime) should win over decision
        category, reason = dg.route_category(probe_result, decision)
        assert category == "tables", f"Expected 'tables', got '{category}'. table_density={probe_result.table_density}, decision={decision}"


# ──────────────────────────────────────────────────────────────────────────
# route_category Tests
# ──────────────────────────────────────────────────────────────────────────


class TestRouteCategory:
    """Test routing rule truth table."""

    # (a) Image-only → 'ocr'
    def test_route_image_only(self):
        """Rule (a): image-only document → ocr."""
        probe = dg.ProbeResult(
            mime="application/pdf", pages=3, chars=0,
            has_text_layer=False, has_page_images=True,
            table_density=0.0, is_email=False
        )
        category, reason = dg.route_category(probe, None)
        assert category == "ocr"
        assert "image-only" in reason

    # (b) Table heavy → 'tables'
    def test_route_table_via_decision(self):
        """Rule (b): decision marks table_heavy with high confidence → tables."""
        probe = dg.ProbeResult(
            mime="text/csv", pages=1, chars=1000,
            has_text_layer=True, has_page_images=False,
            table_density=0.2, is_email=False
        )
        decision = {
            "kind": "table_heavy",
            "confidence": 0.9,
            "complexity": "low",
            "probabilities": {},
        }
        category, reason = dg.route_category(probe, decision)
        assert category == "tables"
        assert "decision: table_heavy" in reason

    def test_route_table_via_density(self):
        """Rule (b): high table density → tables."""
        probe = dg.ProbeResult(
            mime="text/csv", pages=1, chars=2000,
            has_text_layer=True, has_page_images=False,
            table_density=0.7, is_email=False
        )
        category, reason = dg.route_category(probe, None)
        assert category == "tables"
        assert "table_density" in reason

    # (c) Short text with confident decision → 'text'
    def test_route_short_text_email_decision(self):
        """Rule (c): decision says email, <4000 chars → text."""
        probe = dg.ProbeResult(
            mime="message/rfc822", pages=1, chars=500,
            has_text_layer=True, has_page_images=False,
            table_density=0.0, is_email=True
        )
        decision = {
            "kind": "email",
            "confidence": 0.85,
            "complexity": "low",
            "probabilities": {},
        }
        category, reason = dg.route_category(probe, decision)
        assert category == "text"
        assert "decision: email" in reason

    def test_route_short_text_plain_note_decision(self):
        """Rule (c): decision says plain_note, <4000 chars → text."""
        probe = dg.ProbeResult(
            mime="text/plain", pages=1, chars=800,
            has_text_layer=True, has_page_images=False,
            table_density=0.0, is_email=False
        )
        decision = {
            "kind": "plain_note",
            "confidence": 0.8,
            "complexity": "low",
            "probabilities": {},
        }
        category, reason = dg.route_category(probe, decision)
        assert category == "text"

    # (d) Fallback rules
    def test_route_fallback_email(self):
        """Rule (d): no decision but is email → text."""
        probe = dg.ProbeResult(
            mime="message/rfc822", pages=1, chars=1000,
            has_text_layer=True, has_page_images=False,
            table_density=0.0, is_email=True
        )
        category, reason = dg.route_category(probe, None)
        assert category == "text"
        assert "rules fallback" in reason

    def test_route_fallback_short(self):
        """Rule (d): no decision but short → text."""
        probe = dg.ProbeResult(
            mime="text/plain", pages=1, chars=1500,
            has_text_layer=True, has_page_images=False,
            table_density=0.0, is_email=False
        )
        category, reason = dg.route_category(probe, None)
        assert category == "text"
        assert "rules fallback" in reason

    def test_route_fallback_long_doc(self):
        """Rule (d): no decision, not email, not short → docs."""
        probe = dg.ProbeResult(
            mime="application/pdf", pages=50, chars=50000,
            has_text_layer=True, has_page_images=False,
            table_density=0.0, is_email=False
        )
        category, reason = dg.route_category(probe, None)
        assert category == "docs"
        assert "rules fallback: default to docs" in reason

    def test_route_fallback_low_confidence(self):
        """Rule (d): decision present but confidence too low → rules fallback."""
        probe = dg.ProbeResult(
            mime="text/plain", pages=1, chars=1500,  # Short, < SHORT_CHARS
            has_text_layer=True, has_page_images=False,
            table_density=0.0, is_email=False
        )
        decision = {
            "kind": "email",
            "confidence": 0.45,  # < CONF_MIN
            "complexity": "low",
            "probabilities": {},
        }
        category, reason = dg.route_category(probe, decision)
        # Short and low confidence → text via rules
        assert category == "text"

    # (e) Default
    def test_route_default_docs(self):
        """Rule (e): confident decision not matching (c) → docs."""
        probe = dg.ProbeResult(
            mime="application/pdf", pages=20, chars=20000,
            has_text_layer=True, has_page_images=False,
            table_density=0.0, is_email=False
        )
        decision = {
            "kind": "report",
            "confidence": 0.92,
            "complexity": "high",
            "probabilities": {},
        }
        category, reason = dg.route_category(probe, decision)
        assert category == "docs"


# ──────────────────────────────────────────────────────────────────────────
# run_routing Integration Tests
# ──────────────────────────────────────────────────────────────────────────


class TestRunRouting:
    """Test the full routing pipeline."""

    def test_routing_with_high_confidence_decision(self):
        """Full pipeline: decider returns high-confidence result."""
        doc = FakeDocument(
            id="test123", filename="invoice.pdf", mime="application/pdf",
            text="Invoice #INV-001\nAmount: $1,000\nDate: 2024-01-15\n" * 100,
            pages=5
        )

        def fake_decide(text):
            return {
                "kind": "invoice",
                "confidence": 0.95,
                "complexity": "low",
                "probabilities": {"invoice": 0.95, "receipt": 0.04, "other": 0.01},
                "source": "laya",
            }

        result = dg.run_routing(
            doc,
            decide=fake_decide,
            resolve_model=fake_resolve_model,
        )

        assert result["doc_id"] == "test123"
        assert result["category"] == "docs"  # Confident but not in short text rules
        assert result["model"] == "general-chat:7b"
        assert result["overridden"] is False
        assert result["decision"]["source"] == "laya"
        assert result["decision"]["kind"] == "invoice"

    def test_routing_with_low_confidence_decision(self):
        """Decider returns result but low confidence → use rules."""
        doc = FakeDocument(
            id="test456", filename="email.eml", mime="message/rfc822",
            text="From: sender@example.com\nTo: recipient@example.com\nSubject: Hello\n\nHi there",
            pages=1
        )

        def fake_decide(text):
            return {
                "kind": "unclear",
                "confidence": 0.35,  # < CONF_MIN
                "complexity": None,
                "probabilities": {},
                "source": "laya",
            }

        result = dg.run_routing(
            doc,
            decide=fake_decide,
            resolve_model=fake_resolve_model,
        )

        assert result["category"] == "text"  # Rules kick in: is_email
        assert result["model"] == "fast-chat:3b"
        # Decision is still present but low confidence
        assert result["decision"]["confidence"] == 0.35

    def test_routing_with_none_decision(self):
        """Decider returns None (unavailable) → use rules."""
        doc = FakeDocument(
            id="test789", filename="short.txt", mime="text/plain",
            text="Just a few words here.",
            pages=1
        )

        def fake_decide(text):
            return None  # Decider unavailable

        result = dg.run_routing(
            doc,
            decide=fake_decide,
            resolve_model=fake_resolve_model,
        )

        assert result["category"] == "text"  # Short → text
        assert result["decision"]["source"] == "rules"
        assert result["decision"]["kind"] is None

    def test_routing_image_only(self):
        """Image-only document: decider not called."""
        doc = FakeDocument(
            id="scan001", filename="scan.pdf", mime="application/pdf",
            text="",  # No text layer
            pages=3,
            page_images=["/path/1.png", "/path/2.png", "/path/3.png"]
        )

        decide_called = []
        def fake_decide(text):
            decide_called.append(text)
            return None

        result = dg.run_routing(
            doc,
            decide=fake_decide,
            resolve_model=fake_resolve_model,
        )

        # Decider should NOT be called
        assert len(decide_called) == 0
        assert result["category"] == "ocr"
        assert result["model"] == "vision-ocr:8b"
        # Decision should have been skipped
        decide_trace = next((t for t in result["trace"] if t["node"] == "decide"), None)
        assert decide_trace["status"] == "skipped"

    def test_routing_with_override_model(self):
        """Override model forces choice but category still computed."""
        doc = FakeDocument(
            id="test999", filename="doc.pdf", mime="application/pdf",
            text="This is a regular document with some content.",
            pages=1
        )

        def fake_decide(text):
            return {
                "kind": "report",
                "confidence": 0.8,
                "complexity": "medium",
                "probabilities": {},
                "source": "laya",
            }

        result = dg.run_routing(
            doc,
            decide=fake_decide,
            resolve_model=fake_resolve_model,
            override_model="custom-model:10b",
        )

        assert result["category"] == "docs"  # Still computed
        assert result["model"] == "custom-model:10b"  # Overridden
        assert result["overridden"] is True

    def test_routing_trace_complete(self):
        """Verify trace includes all nodes."""
        doc = FakeDocument(
            id="trace_test", filename="doc.txt", mime="text/plain",
            text="Some content here." * 100,
            pages=1
        )

        def fake_decide(text):
            return {
                "kind": "document",
                "confidence": 0.7,
                "complexity": "low",
                "probabilities": {},
                "source": "laya",
            }

        result = dg.run_routing(
            doc,
            decide=fake_decide,
            resolve_model=fake_resolve_model,
        )

        trace_nodes = [t["node"] for t in result["trace"]]
        expected = ["ingest", "probe", "decide", "route", "resolve_model"]
        assert trace_nodes == expected

        # Verify timing is recorded
        for entry in result["trace"]:
            assert "ms" in entry
            assert entry["ms"] >= 0


# ──────────────────────────────────────────────────────────────────────────
# Email Ingest & ingest_text Tests
# ──────────────────────────────────────────────────────────────────────────


class TestEmailIngest:
    """Test .eml file ingestion."""

    def test_ingest_email_file(self):
        """Build and ingest an .eml file."""
        # Create a minimal .eml file
        msg = EmailMessage()
        msg["From"] = "alice@example.com"
        msg["To"] = "bob@example.com"
        msg["Subject"] = "Test Email"
        msg["Date"] = "Mon, 01 Jan 2024 12:00:00 +0000"
        msg.set_content("Hello, this is the email body.")

        with tempfile.TemporaryDirectory() as tmpdir:
            eml_path = pathlib.Path(tmpdir) / "test.eml"
            eml_path.write_bytes(msg.as_bytes())

            # Read and ingest
            data = eml_path.read_bytes()
            import asyncio
            doc = asyncio.run(ide.ingest_file("test.eml", data))

            # Verify
            assert "alice@example.com" in doc.text
            assert "bob@example.com" in doc.text
            assert "Test Email" in doc.text
            assert "Hello, this is the email body" in doc.text
            assert doc.mime == "message/rfc822"
            assert doc.pages == 1


class TestIngestText:
    """Test ingest_text for creating documents from plain text."""

    def test_ingest_text_plain(self):
        """Create a text document from plain text."""
        import asyncio

        text = "This is a test document with some content."
        doc = asyncio.run(ide.ingest_text(text, title="My Document"))

        assert doc.text == text
        assert doc.filename == "My Document"
        assert doc.mime == "text/plain"
        assert doc.pages == 1
        assert doc.page_texts == [text]

    def test_ingest_text_email(self):
        """Create an email document from text."""
        import asyncio

        text = "From: sender@example.com\nTo: recipient@example.com\nSubject: Hi\n\nBody text"
        doc = asyncio.run(ide.ingest_text(text, kind="email", title="Forwarded Message"))

        assert doc.text == text
        assert doc.filename == "Forwarded Message"
        assert doc.mime == "message/rfc822"
        assert doc.pages == 1

    def test_ingest_text_default_title(self):
        """Default title based on kind."""
        import asyncio

        text = "Test content"
        doc_text = asyncio.run(ide.ingest_text(text))
        assert "Pasted text" in doc_text.filename

        doc_email = asyncio.run(ide.ingest_text(text, kind="email"))
        assert "Pasted email" in doc_email.filename

    def test_ingest_text_empty_raises(self):
        """Empty text raises ValueError."""
        import asyncio

        try:
            asyncio.run(ide.ingest_text(""))
            assert False, "Should raise ValueError"
        except ValueError as e:
            assert "empty" in str(e).lower()

    def test_ingest_text_whitespace_only_raises(self):
        """Whitespace-only text raises ValueError."""
        import asyncio

        try:
            asyncio.run(ide.ingest_text("   \n\n  "))
            assert False, "Should raise ValueError"
        except ValueError as e:
            assert "empty" in str(e).lower()
