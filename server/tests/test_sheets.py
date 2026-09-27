"""
Spreadsheet support: extraction (sheets.py / idp_engine), the sheet preview
handler and the doc agent's deterministic "Table query" step.

All fixtures are generated in tmp_path; ingest tests write only to the
isolated STORAGE_DIR set up by conftest.py.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

pd = pytest.importorskip("pandas")
pytest.importorskip("openpyxl")

import doc_agent as _agent
import idp_engine as _idp
import sheets as _sheets


# ── Fixtures ────────────────────────────────────────────────────────────────

def make_workbook(path: Path) -> Path:
    """Sales (Region, Q1..Q4, formula Total), Staff (merged title banner,
    header on row 2, dates, a vertically merged Department cell) and a hidden
    sheet."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1", "Q2", "Q3", "Q4", "Total"])
    data = [("East", 10, 20, 30, 40), ("West", 5, 5, 5, 5), ("North", 1, 2, 3, 4), ("South", 7, 8, 9, 10)]
    for i, (region, *q) in enumerate(data, start=2):
        ws.append([region, *q, f"=SUM(B{i}:E{i})"])
    ws.append([None, None, None, None, None, None])       # fully-empty row → dropped

    st = wb.create_sheet("Staff")
    st.append(["Staff list 2024"])
    st.merge_cells("A1:D1")                                 # merged title banner
    st.append(["Name", "Department", "Salary", "Start date"])
    st.append(["Ann", "Sales", 50000, dt.datetime(2020, 1, 15)])
    st.append(["Bob", "IT", 60000, dt.date(2021, 3, 1)])
    st.append(["Cid", None, 70000, dt.datetime(2019, 7, 1)])
    st.merge_cells("B4:B5")                                 # Bob + Cid both IT

    hidden = wb.create_sheet("Secret")
    hidden.append(["Key", "Value"])
    hidden.append(["pin", 1234])
    hidden.sheet_state = "hidden"
    wb.save(path)
    return path


@pytest.fixture
def xlsx_path(tmp_path: Path) -> Path:
    return make_workbook(tmp_path / "report.xlsx")


def _sheet(sheets: list[dict], name: str) -> dict:
    return next(s for s in sheets if s["name"] == name)


# ── Extraction ──────────────────────────────────────────────────────────────

class TestXlsxExtraction:
    def test_sheet_names_and_hidden_flag(self, xlsx_path):
        sheets = _sheets.parse_workbook(xlsx_path)
        assert [s["name"] for s in sheets] == ["Sales", "Staff", "Secret"]
        assert _sheet(sheets, "Secret")["hidden"] is True
        assert not _sheet(sheets, "Sales")["hidden"]

    def test_header_detection_and_counts(self, xlsx_path):
        sheets = _sheets.parse_workbook(xlsx_path)
        sales = _sheet(sheets, "Sales")
        assert list(sales["df"].columns) == ["Region", "Q1", "Q2", "Q3", "Q4", "Total"]
        assert len(sales["df"]) == 4                        # empty row dropped
        assert sales["header_row"] == 1
        staff = _sheet(sheets, "Staff")
        # merged title banner is not taken for the header
        assert list(staff["df"].columns) == ["Name", "Department", "Salary", "Start date"]
        assert staff["header_row"] == 2
        assert staff["preamble"] == ["Staff list 2024"]
        assert len(staff["df"]) == 3

    def test_merged_cells_forward_filled(self, xlsx_path):
        staff = _sheet(_sheets.parse_workbook(xlsx_path), "Staff")
        assert staff["df"]["Department"].tolist() == ["Sales", "IT", "IT"]

    def test_dates_iso_and_numbers_numeric(self, xlsx_path):
        staff = _sheet(_sheets.parse_workbook(xlsx_path), "Staff")
        assert staff["df"]["Start date"].tolist() == ["2020-01-15", "2021-03-01", "2019-07-01"]
        assert staff["dtypes"]["Start date"] == "date"
        assert staff["dtypes"]["Salary"] == "number"
        assert pd.api.types.is_numeric_dtype(staff["df"]["Salary"])

    def test_formula_column_recorded_and_evaluated(self, xlsx_path):
        # openpyxl-written workbooks have no cached formula values: the simple
        # SUM is evaluated by the safe fallback evaluator.
        sales = _sheet(_sheets.parse_workbook(xlsx_path), "Sales")
        assert sales["formula_columns"] == ["Total"]
        assert sales["df"]["Total"].tolist() == [100, 20, 10, 34]

    def test_pages_and_meta(self, xlsx_path, tmp_path):
        cache = tmp_path / _sheets.CACHE_NAME
        pages, meta = _sheets.extract(xlsx_path, cache)
        assert len(pages) == 2                                # hidden sheet left out
        assert pages[0].startswith("### Sheet: Sales (4 rows × 6 columns)\n")
        assert "| Region | Q1 | Q2 | Q3 | Q4 | Total |" in pages[0]
        assert "| East | 10 | 20 | 30 | 40 | 100 |" in pages[0]
        assert "Q1 — sum 23, mean 5.75, min 1, max 10" in pages[0]
        assert "Formula columns: Total" in pages[0]
        assert "Secret" not in "\n".join(pages)
        names = [s["name"] for s in meta["sheets"]]
        assert names == ["Sales", "Staff", "Secret"]
        secret = meta["sheets"][2]
        assert secret["hidden"] is True
        sales = meta["sheets"][0]
        assert sales["rows"] == 4 and sales["cols"] == 6
        total_col = next(c for c in sales["columns"] if c["name"] == "Total")
        assert total_col["formula"] is True and total_col["dtype"] == "number"
        assert sales["columns"][0] == {"name": "Region", "dtype": "text", "non_null": 4,
                                       "sample": ["East", "West", "North"]}
        assert meta["hidden_sheets"] == ["Secret"]
        assert meta["sheet_count"] == 2
        assert cache.exists()

    def test_row_cap_in_text(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Big"
        ws.append(["Id", "Amount"])
        for i in range(1, 451):
            ws.append([i, i * 2])
        p = tmp_path / "big.xlsx"
        wb.save(p)
        pages, meta = _sheets.extract(p)
        assert "### Sheet: Big (450 rows × 2 columns)" in pages[0]
        assert "… 150 more rows" in pages[0]
        assert "| 300 | 600 |" in pages[0]
        assert "| 301 | 602 |" not in pages[0]
        assert "Amount — sum 202950" in pages[0]         # stats over ALL rows
        assert meta["sheets"][0]["rows"] == 450

    def test_no_header_generates_column_letters(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append([None, 1, 2])
        ws.append([None, 3, 4])
        p = tmp_path / "nums.xlsx"
        wb.save(p)
        sh = _sheets.parse_workbook(p)[0]
        assert list(sh["df"].columns) == ["Column B", "Column C"]
        assert sh["header_row"] is None
        assert len(sh["df"]) == 2


class TestOtherFormats:
    def test_ods(self, tmp_path):
        pytest.importorskip("odf")
        p = tmp_path / "book.ods"
        with pd.ExcelWriter(p, engine="odf") as xw:
            pd.DataFrame({"City": ["Paris", "Rome"], "Pop": [2.1, 2.8]}).to_excel(xw, sheet_name="Cities", index=False)
            pd.DataFrame({"K": ["a"], "V": [1]}).to_excel(xw, sheet_name="Other", index=False)
        sheets = _sheets.parse_workbook(p)
        assert [s["name"] for s in sheets] == ["Cities", "Other"]
        cities = sheets[0]
        assert list(cities["df"].columns) == ["City", "Pop"]
        assert cities["df"]["Pop"].tolist() == [2.1, 2.8]
        pages, meta = _sheets.extract(p)
        assert pages[0].startswith("### Sheet: Cities (2 rows × 2 columns)")
        assert meta["sheet_format"] == "ods"

    def test_csv_semicolon_and_numbers(self, tmp_path):
        p = tmp_path / "sales.csv"
        p.write_text("Region;Amount;Date\nEast;1,200.5;2024-01-02\nWest;300;2024-02-03\n", encoding="utf-8")
        sh = _sheets.parse_workbook(p)[0]
        assert sh["name"] == "sales"
        assert list(sh["df"].columns) == ["Region", "Amount", "Date"]
        assert sh["df"]["Amount"].tolist() == [1200.5, 300]
        assert sh["dtypes"]["Amount"] == "number"
        assert sh["dtypes"]["Date"] == "date"

    def test_tsv(self, tmp_path):
        p = tmp_path / "t.tsv"
        p.write_text("A\tB\nx\t1\ny\t2\n", encoding="utf-8")
        sh = _sheets.parse_workbook(p)[0]
        assert list(sh["df"].columns) == ["A", "B"]
        assert sh["df"]["B"].sum() == 3

    def test_xls_dispatch_uses_xlrd_loader(self, tmp_path, monkeypatch):
        """.xls goes through the xlrd loader (no .xls writer is available to
        build a real fixture, so the loader is stubbed with a raw grid)."""
        called = {}

        def fake_load(path):
            called["path"] = path
            return [_sheets.RawSheet("Legacy", [["Item", "Qty"], ["Bolt", 3], ["Nut", 4]]),
                    _sheets.RawSheet("Old", [["x"]], hidden=True)]

        monkeypatch.setattr(_sheets, "_load_xls", fake_load)
        p = tmp_path / "legacy.xls"
        p.write_bytes(b"\xd0\xcf\x11\xe0 not really")
        pages, meta = _sheets.extract(p)
        assert called["path"] == p
        assert pages[0].startswith("### Sheet: Legacy (2 rows × 2 columns)")
        assert meta["hidden_sheets"] == ["Old"]
        assert meta["sheet_format"] == "xls"

    def test_real_xls_with_xlwt(self, tmp_path):
        xlwt = pytest.importorskip(
            "xlwt", reason="no .xls writer (xlwt) installed and no LibreOffice to create a real .xls fixture",
        )
        pytest.importorskip("xlrd")
        wb = xlwt.Workbook()
        ws = wb.add_sheet("Legacy")
        for r, row in enumerate([["Item", "Qty"], ["Bolt", 3], ["Nut", 4]]):
            for c, v in enumerate(row):
                ws.write(r, c, v)
        p = tmp_path / "legacy.xls"
        wb.save(str(p))
        sh = _sheets.parse_workbook(p)[0]
        assert list(sh["df"].columns) == ["Item", "Qty"]
        assert sh["df"]["Qty"].tolist() == [3, 4]


# ── idp_engine integration ──────────────────────────────────────────────────

class TestIdpIntegration:
    def test_mime_and_dispatch_lists(self):
        assert _idp._detect_mime("a.xls") == "application/vnd.ms-excel"
        assert _idp._detect_mime("a.xlsm") == "application/vnd.ms-excel.sheet.macroEnabled.12"
        assert _idp._detect_mime("a.ods") == "application/vnd.oasis.opendocument.spreadsheet"
        assert _idp._detect_mime("a.tsv") == "text/tab-separated-values"
        for ext in (".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".tsv"):
            assert ext in _idp.REEXTRACTABLE_EXTS
            assert _idp.is_sheet_file("x" + ext)
        assert not _idp.is_sheet_file("x.pdf", "application/pdf")

    def test_ingest_preview_and_cache_regeneration(self, xlsx_path):
        doc = asyncio.run(_idp.ingest_file("report.xlsx", xlsx_path.read_bytes()))
        try:
            assert doc.pages == 2
            assert isinstance(doc.meta["sheets"], list)
            assert [s["name"] for s in doc.meta["sheets"]] == ["Sales", "Staff", "Secret"]
            cache = _idp.sheet_cache_path(doc)
            assert cache.exists()

            prev = _idp.sheet_preview(doc.id)
            assert prev["sheet"] == "Sales"
            assert prev["columns"] == ["Region", "Q1", "Q2", "Q3", "Q4", "Total"]
            assert prev["dtypes"][1] == "number"
            assert prev["total"] == 4 and prev["offset"] == 0
            assert prev["rows"][0] == ["East", 10, 20, 30, 40, 100]
            # hidden sheets are not listed by default (only counted)
            assert [s["name"] for s in prev["sheets"]] == ["Sales", "Staff"]
            assert prev["hidden_sheets"] == 1
            with_hidden = _idp.sheet_preview(doc.id, include_hidden=True)
            assert [s["name"] for s in with_hidden["sheets"]] == ["Sales", "Staff", "Secret"]
            assert with_hidden["sheets"][2]["hidden"] is True

            page = _idp.sheet_preview(doc.id, "staff", offset=1, limit=1)
            assert page["sheet"] == "Staff"
            assert page["rows"] == [["Bob", "IT", 60000, "2021-03-01"]]
            assert page["offset"] == 1 and page["total"] == 3

            # cache deleted → regenerated from the raw file
            cache.unlink()
            _sheets.forget(cache)
            again = _idp.sheet_preview(doc.id, "Sales")
            assert again["total"] == 4
            assert cache.exists()

            with pytest.raises(LookupError):
                _idp.sheet_preview(doc.id, "Nope")
            with pytest.raises(LookupError):
                _idp.sheet_preview("missing-doc")
        finally:
            _idp.delete_document(doc.id)

    def test_preview_rejects_non_spreadsheet(self):
        doc = asyncio.run(_idp.ingest_text("hello world", title="note"))
        try:
            with pytest.raises(ValueError):
                _idp.sheet_preview(doc.id)
        finally:
            _idp.delete_document(doc.id)


# ── Doc agent: Table query step ─────────────────────────────────────────────

def _sheet_doc(path: Path, doc_id: str = "sheet1"):
    sheets = _sheets.parse_workbook(path)
    pages = _sheets.render_pages(sheets)
    meta = _sheets.build_meta(sheets, path.suffix.lstrip("."))
    doc = SimpleNamespace(
        id=doc_id, filename=path.name,
        mime=_idp._detect_mime(path.name), size=1000, uploaded_at=0.0,
        pages=len(pages), text="\n\n".join(pages), page_texts=pages,
        page_images=[], meta=meta,
    )
    return doc, sheets


def _hooks(docs: dict, frames: dict, replies: list[str], prompts: list[str], kwargs_seen: list[dict]):
    """Minimal AgentHooks: the first stream_llm calls return `replies` (query
    specs), later calls return a normal answer."""
    replies = list(replies)

    async def stream_llm(model, prompt, think=False, temperature=None, num_predict=1536):
        prompts.append(prompt)
        kwargs_seen.append({"temperature": temperature, "num_predict": num_predict})
        if "JSON:" in prompt[-10:] and replies:
            text = replies.pop(0)
        else:
            text = "The answer is in the computed result."
        yield {"content": text}
        yield {"done": True, "stats": {}}

    async def _const(*a, **k):
        return "text-model"

    async def _info(*a, **k):
        return {"name": "text-model"}

    async def _none(*a, **k):
        return None

    async def _empty(*a, **k):
        return []

    return _agent.AgentHooks(
        get_doc=lambda i: docs.get(i),
        laya_intent=lambda m, f: None,
        laya_role=lambda m, f: None,
        laya_doc_kind=lambda t: None,
        laya_info=lambda: {"name": "Laya", "available": False, "device": None},
        page_image_paths=lambda d, p=None, dpi=120: [],
        resolve_model=_const,
        resolve_text_model=_const,
        pick_vision_model=_none,
        vision_candidates=_empty,
        model_info=_info,
        run_ocr=_const,
        stream_llm=stream_llm,
        vision_call=_const,
        mark_vision_failed=lambda m, e: None,
        sheet_frames=lambda d: frames.get(d.id),
        now_ms=lambda: 1_000_000,
    )


def _run(hooks, message: str, doc_id: str = "sheet1") -> list[dict]:
    async def go():
        return [ev async for ev in _agent.run_agent(
            {"message": message, "attachments": [{"doc_id": doc_id}]}, hooks,
        )]
    return asyncio.run(go())


def _tiles(events: list[dict]) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for ev in events:
        if "tile" in ev:
            out[ev["tile"]["id"]] = ev["tile"]
    return out


class TestAgentTableQuery:
    def test_table_question_detection(self):
        assert "average" in _agent.table_question_hits("What is the average salary per department?")
        assert "hoogste" in _agent.table_question_hits("Welke regio heeft de hoogste omzet?")
        assert _agent.table_question_hits("Summarize this document") == []

    def test_across_is_local_for_spreadsheets(self):
        msg = "What is the total across Q1..Q4 for East?"
        assert _agent.rules_intent(msg, [{"name": "r.xlsx", "kind": "xlsx"}])[0] != "graph_query"
        assert _agent.rules_intent(msg, [{"name": "a.pdf", "kind": "pdf"}])[0] == "graph_query"

    def test_row_total_question_end_to_end(self, xlsx_path):
        doc, sheets = _sheet_doc(xlsx_path)
        prompts, kw = [], []
        spec = {"row_total": {"columns": "Q1..Q4", "as": "Total"},
                "filters": [{"column": "Region", "op": "==", "value": "East"}], "select": ["Region"]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(spec)], prompts, kw)
        events = _run(hooks, "What is the total across Q1..Q4 for East?")
        tq = _tiles(events)["table-query"]
        assert tq["status"] == "done"
        assert "| East | 100 |" in prompts[-1]

    def test_query_tile_and_answer_prompt(self, xlsx_path):
        doc, sheets = _sheet_doc(xlsx_path)
        prompts, kw = [], []
        spec = {"sheet": "Staff", "group_by": ["Department"],
                "aggregate": [{"column": "Salary", "fn": "mean"}],
                "sort": [{"column": "Salary", "desc": True}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(spec)], prompts, kw)
        events = _run(hooks, "What is the average salary per department?")
        assert not [e for e in events if "error" in e], events
        tiles = _tiles(events)
        tq = tiles["table-query"]
        assert tq["kind"] == "table" and tq["status"] == "done"
        assert tq["model"] == "text-model"
        assert tq["detail"] == "2 rows · sheet Staff"
        assert tq["items"][0]["kind"] == "query"
        assert "mean(Salary) by Department" in tq["items"][0]["label"]
        assert tq["items"][0]["columns"] == ["Department", "mean(Salary)"]
        results = [i for i in tq["items"] if i["kind"] == "result"]
        assert results[0]["row"] == ["IT", 65000] and results[1]["row"] == ["Sales", 50000]
        assert results[0]["label"] == "Department: IT · mean(Salary): 65000"
        # tile order: table query before the answer
        order = [e["tile"]["id"] for e in events if "tile" in e]
        assert order.index("table-query") < order.index("answer")
        # the spec prompt carries the schema; the answer prompt the computed result
        assert "Sheet \"Staff\"" in prompts[0] and "Salary (number)" in prompts[0]
        answer_prompt = prompts[-1]
        assert "COMPUTED RESULT" in answer_prompt
        assert "| IT | 65000 |" in answer_prompt
        assert "do not recalculate" in answer_prompt
        # low temperature for the spec call when the hook supports it
        assert kw[0]["temperature"] == 0.0
        laya = tiles["laya"]
        assert any(d["id"] == "table_question" for d in laya["decisions"])

    def test_retry_after_validation_error(self, xlsx_path):
        doc, sheets = _sheet_doc(xlsx_path)
        prompts, kw = [], []
        bad = {"sheet": "Sales", "aggregate": [{"column": "Revenue", "fn": "sum"}]}
        good = {"sheet": "Sales", "aggregate": [{"column": "Q1", "fn": "sum"}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(bad), json.dumps(good)], prompts, kw)
        events = _run(hooks, "What is the total of Q1?")
        tq = _tiles(events)["table-query"]
        assert tq["status"] == "done"
        assert "Unknown column 'Revenue'" in prompts[1]
        assert "Available columns: Region, Q1" in prompts[1]
        assert "| 23 |" in prompts[-1]

    def test_failed_query_falls_back(self, xlsx_path):
        doc, sheets = _sheet_doc(xlsx_path)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, ["no idea", "still no json"], prompts, kw)
        events = _run(hooks, "What is the highest total?")
        tiles = _tiles(events)
        tq = tiles["table-query"]
        assert tq["status"] == "error"
        assert tiles["answer"]["status"] == "done"
        answer_prompt = prompts[-1]
        assert "COMPUTED RESULT" not in answer_prompt
        # per-column stats still reach the model
        assert "Stats (all rows): Q1 — sum 23" in answer_prompt

    def test_large_sheet_prompt_never_contains_all_rows(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Orders"
        ws.append(["Order", "Customer", "Amount"])
        for i in range(1, 1001):
            ws.append([f"ORD-{i:05d}", f"Customer {i % 37}", i])
        p = tmp_path / "orders.xlsx"
        wb.save(p)
        doc, sheets = _sheet_doc(p)
        prompts, kw = [], []
        spec = {"aggregate": [{"column": "Amount", "fn": "sum"}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(spec)], prompts, kw)
        events = _run(hooks, "What is the total amount?")
        assert _tiles(events)["table-query"]["status"] == "done"
        answer_prompt = prompts[-1]
        assert "| 500500 |" in answer_prompt
        assert "ORD-00001" in answer_prompt           # sample rows
        assert "ORD-00500" not in answer_prompt       # …but never the whole sheet
        assert "ORD-01000" not in answer_prompt
        for pr in prompts:
            assert "ORD-00999" not in pr

    def test_large_sheet_fallback_also_capped(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Order", "Amount"])
        for i in range(1, 801):
            ws.append([f"ORD-{i:05d}", i])
        p = tmp_path / "o.xlsx"
        wb.save(p)
        doc, sheets = _sheet_doc(p)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, ["??", "??"], prompts, kw)
        _run(hooks, "Which order has the highest amount?")
        assert "ORD-00800" not in prompts[-1]
        assert "Amount — sum 320400" in prompts[-1]

    def test_summarize_does_not_run_table_query(self, xlsx_path):
        doc, sheets = _sheet_doc(xlsx_path)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [], prompts, kw)
        events = _run(hooks, "Summarize this spreadsheet")
        assert "table-query" not in _tiles(events)
        assert len(prompts) == 1

    def test_non_sheet_doc_unaffected(self):
        doc = SimpleNamespace(
            id="d1", filename="note.pdf", mime="application/pdf", size=1, uploaded_at=0.0,
            pages=1, text="Invoice total 100 EUR", page_texts=["Invoice total 100 EUR"],
            page_images=[], meta={},
        )
        prompts, kw = [], []
        hooks = _hooks({"d1": doc}, {}, [], prompts, kw)
        events = _run(hooks, "What is the total?", doc_id="d1")
        assert "table-query" not in _tiles(events)

    def test_extract_tile_detail_for_spreadsheet(self, xlsx_path):
        doc, sheets = _sheet_doc(xlsx_path)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [], prompts, kw)
        events = _run(hooks, "Summarize this spreadsheet")
        assert _tiles(events)[f"extract-{doc.id}"]["detail"] == "Spreadsheet · 2 sheets · 7 rows"

    def test_mime_to_kind_spreadsheets(self):
        assert _agent._mime_to_kind("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet") == "xlsx"
        assert _agent._mime_to_kind("application/vnd.ms-excel") == "xlsx"
        assert _agent._mime_to_kind("application/vnd.oasis.opendocument.spreadsheet") == "xlsx"
        assert _agent._mime_to_kind("application/vnd.openxmlformats-officedocument.wordprocessingml.document") == "docx"
