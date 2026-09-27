"""
Spreadsheet fixes from the live test matrix (company.xlsx / inventory.ods /
big.csv, regenerated here in tmp_path from the matrix generator's logic):

1. sort keys resolve before the projection (select + sort on another column)
2. "derive" for products / ratios; row_total-for-a-product is rejected
3. hidden sheets never leak (query, meta, preview, schema, prompts)
4. follow-up questions get the previous turns in the spec prompt
5. row_total reuses an existing matching Total column
6. tile detail when the prompt got the schema instead of the sheet
7. legacy dict-style meta["sheets"] is normalised
"""

from __future__ import annotations

import asyncio
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent))   # reuse the test_sheets helpers

import pytest

pd = pytest.importorskip("pandas")
pytest.importorskip("openpyxl")

import doc_agent as _agent
import idp_engine as _idp
import sheets as _sheets
import table_query as tq

from test_sheets import _hooks, _sheet_doc, _tiles

SECRET = "launch code 7741-ZEBRA"
STAFF = [
    ("Alice Jansen", "Engineering", 85000, dt.date(2019, 3, 1)),
    ("Bob de Vries", "Engineering", 92000, dt.date(2021, 6, 15)),
    ("Carla Smit", "Engineering", 78000, dt.date(2023, 1, 10)),
    ("Daan Bakker", "Sales", 61000, dt.date(2020, 9, 1)),
    ("Eva Mulder", "Sales", 67000, dt.date(2022, 11, 20)),
    ("Frank Visser", "Sales", 58000, dt.date(2024, 2, 5)),
    ("Gina Bos", "HR", 55000, dt.date(2019, 7, 22)),
    ("Hugo Peters", "HR", 59000, dt.date(2025, 4, 1)),
    ("Iris Dekker", "Finance", 98000, dt.date(2020, 1, 13)),
    ("Jan Hendriks", "Finance", 73000, dt.date(2023, 8, 30)),
    ("Kim Vos", "Finance", 81000, dt.date(2022, 5, 16)),
    ("Lars Meijer", "Engineering", 105000, dt.date(2025, 1, 6)),
]
SALES = [("North", 120, 135, 150, 160), ("South", 90, 95, 80, 110),
         ("East", 200, 210, 190, 220), ("West", 60, 75, 85, 95)]
INVENTORY = {
    "Item": ["Widget", "Gadget", "Bolt", "Nut", "Screwdriver", "Hammer", "Drill", "Saw"],
    "Category": ["Parts", "Parts", "Hardware", "Hardware", "Tools", "Tools", "Tools", "Tools"],
    "Stock": [120, 45, 500, 800, 30, 25, 10, 15],
    "Unit price": [2.5, 12.0, 0.1, 0.05, 8.0, 15.0, 89.0, 22.0],
}
STOCK_VALUE = 2765.0     # ground_truth_sheets.json q8


def make_company(path: Path) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Sales 2025"
    ws.merge_cells("A1:F1")
    ws.append(["Region", "Q1", "Q2", "Q3", "Q4", "Total"])
    for i, (r, *q) in enumerate(SALES, start=3):
        ws.append([r, *q, f"=SUM(B{i}:E{i})"])
    st = wb.create_sheet("Staff")
    st.append(["Name", "Department", "Salary", "Start date"])
    for row in STAFF:
        st.append(list(row))
    st.merge_cells("B5:B6")
    for r in range(2, 14):
        st.cell(r, 4).number_format = "yyyy-mm-dd"
    sec = wb.create_sheet("Secret")
    sec["A1"] = SECRET
    sec.sheet_state = "hidden"
    wb.save(path)
    return path


def make_inventory_xlsx(path: Path) -> Path:
    pd.DataFrame(INVENTORY).to_excel(path, index=False)
    return path


@pytest.fixture
def company(tmp_path: Path) -> Path:
    return make_company(tmp_path / "company.xlsx")


@pytest.fixture
def company_book(company: Path) -> list[dict]:
    return _sheets.parse_workbook(company)


@pytest.fixture
def inventory_book(tmp_path: Path) -> list[dict]:
    return _sheets.parse_workbook(make_inventory_xlsx(tmp_path / "inventory.xlsx"))


def _run_h(hooks, message: str, history=None, doc_id: str = "sheet1") -> list[dict]:
    async def go():
        req = {"message": message, "attachments": [{"doc_id": doc_id}]}
        if history is not None:
            req["history"] = history
        return [ev async for ev in _agent.run_agent(req, hooks)]
    return asyncio.run(go())


# ── 1. sort after select ────────────────────────────────────────────────────

class TestSortBeforeProjection:
    def test_highest_salary_with_select_of_other_columns(self, company_book):
        r = tq.execute(company_book, {"sheet": "Staff", "select": ["Name", "Department"],
                                      "sort": [{"column": "Salary", "desc": True}], "limit": 1})
        assert r["columns"] == ["Name", "Department"]          # sort column dropped again
        assert r["rows"] == [["Lars Meijer", "Engineering"]]
        assert r["spec"]["sort"] == [{"column": "Salary", "desc": True}]

    def test_region_with_highest_q4(self, company_book):
        r = tq.execute(company_book, {"select": ["Region"], "sort": [{"column": "Q4", "desc": True}], "limit": 1})
        assert r["sheet"] == "Sales"
        assert r["columns"] == ["Region"]
        assert r["rows"] == [["East"]]

    def test_limit_applies_after_sort(self, company_book):
        r = tq.execute(company_book, {"sheet": "Staff", "select": ["Name"],
                                      "sort": [{"column": "Salary", "desc": True}], "limit": 3})
        assert r["rows"] == [["Lars Meijer"], ["Iris Dekker"], ["Bob de Vries"]]

    def test_unknown_sort_column_still_errors(self, company_book):
        with pytest.raises(tq.QueryError, match="Unknown sort column 'Bonus'"):
            tq.execute(company_book, {"sheet": "Staff", "select": ["Name"], "sort": [{"column": "Bonus"}]})


# ── 2. derive ───────────────────────────────────────────────────────────────

class TestDerive:
    def test_total_stock_value(self, inventory_book):
        spec = {"derive": [{"as": "Value", "expr": {"op": "mul", "left": "Stock", "right": "Unit price"}}],
                "aggregate": [{"column": "Value", "fn": "sum"}]}
        r = tq.execute(inventory_book, spec)
        assert r["columns"] == ["sum(Value)"]
        assert r["rows"][0][0] == pytest.approx(STOCK_VALUE)
        assert f"{STOCK_VALUE:.0f}" in r["markdown"]
        assert "Value = Stock × Unit price" in tq.spec_summary(r["spec"])

    def test_total_stock_value_from_ods(self, tmp_path):
        pytest.importorskip("odf")
        p = tmp_path / "inventory.ods"
        pd.DataFrame(INVENTORY).to_excel(p, engine="odf", index=False)
        book = _sheets.parse_workbook(p)
        r = tq.execute(book, {"derive": [{"as": "Value", "expr": {"op": "mul", "left": "stock", "right": "unit_price"}}],
                              "aggregate": [{"column": "Value", "fn": "sum"}]})
        assert r["rows"][0][0] == pytest.approx(STOCK_VALUE)

    def test_derive_per_row_sort_group(self, inventory_book):
        spec = {"derive": [{"as": "Value", "expr": {"op": "*", "left": "Stock", "right": "Unit price"}}],
                "select": ["Item"], "sort": [{"column": "Value", "desc": True}], "limit": 1}
        r = tq.execute(inventory_book, spec)
        assert r["columns"] == ["Item", "Value"]          # derived column auto-selected
        assert r["rows"] == [["Drill", 890]]
        g = tq.execute(inventory_book, {
            "derive": [{"as": "Value", "expr": {"op": "mul", "left": "Stock", "right": "Unit price"}}],
            "group_by": ["Category"], "aggregate": [{"column": "Value", "fn": "sum"}]})
        assert dict((k, v) for k, v in g["rows"]) == pytest.approx({"Hardware": 90, "Parts": 840, "Tools": 1835})

    def test_nested_numbers_and_string_form(self, inventory_book):
        pct = {"as": "Pct", "expr": {"op": "mul", "right": 100,
                                     "left": {"op": "div", "left": "Stock", "right": {"op": "add", "left": "Stock", "right": 0}}}}
        r = tq.execute(inventory_book, {"derive": [pct], "aggregate": [{"column": "Pct", "fn": "max"}]})
        assert r["rows"] == [[100]]
        r = tq.execute(inventory_book, {"derive": [{"as": "V", "expr": "Stock * Unit price"}],
                                        "aggregate": [{"column": "V", "fn": "sum"}]})
        assert r["rows"][0][0] == pytest.approx(STOCK_VALUE)
        # a derived column can feed the next one
        r = tq.execute(inventory_book, {"derive": [
            {"as": "V", "expr": {"op": "mul", "left": "Stock", "right": "Unit price"}},
            {"as": "Half", "expr": {"op": "div", "left": "V", "right": 2}}],
            "aggregate": [{"column": "Half", "fn": "sum"}]})
        assert r["rows"][0][0] == pytest.approx(STOCK_VALUE / 2)

    def test_division_by_zero_is_nan(self):
        book = [_sheets.build_sheet(_sheets.RawSheet("T", [["a", "b"], [10, 2], [5, 0]]))]
        r = tq.execute(book, {"derive": [{"as": "r", "expr": {"op": "div", "left": "a", "right": "b"}}], "select": ["a"]})
        assert r["rows"] == [[10, 5], [5, None]]
        r = tq.execute(book, {"derive": [{"as": "r", "expr": {"op": "div", "left": "a", "right": 0}}],
                              "aggregate": [{"column": "r", "fn": "sum"}]})
        assert r["rows"] == [[None]]

    def test_derive_validation(self, inventory_book):
        with pytest.raises(tq.QueryError, match="Allowed ops: mul, div, add, sub"):
            tq.execute(inventory_book, {"derive": [{"as": "V", "expr": {"op": "pow", "left": "Stock", "right": 2}}]})
        deep = {"op": "add", "left": {"op": "add", "left": {"op": "add", "left": {"op": "add", "left": "Stock", "right": 1},
                                                            "right": 1}, "right": 1}, "right": 1}
        with pytest.raises(tq.QueryError, match="nested at most 3"):
            tq.execute(inventory_book, {"derive": [{"as": "V", "expr": deep}]})
        with pytest.raises(tq.QueryError, match="Unknown column 'Weight'"):
            tq.execute(inventory_book, {"derive": [{"as": "V", "expr": {"op": "mul", "left": "Weight", "right": "Stock"}}]})
        with pytest.raises(tq.QueryError, match="already a column"):
            tq.execute(inventory_book, {"derive": [{"as": "Stock", "expr": {"op": "mul", "left": "Stock", "right": 2}}]})
        with pytest.raises(tq.QueryError):
            tq.execute(inventory_book, {"derive": [{"as": "V", "expr": "__import__('os').system('echo hi')"}]})

    def test_row_total_for_product_question_rejected(self):
        spec = {"row_total": {"columns": ["Stock", "Unit price"], "as": "Total Value"},
                "aggregate": [{"column": "Total Value", "fn": "sum"}]}
        with pytest.raises(tq.QueryError, match="use derive for products/ratios"):
            tq.check_spec_for_question("what is the total stock value (stock × unit price)?", spec)
        for q in ("stock x price", "stock times price", "multiply stock by price", "ratio of a to b",
                  "percentage of sales", "share of the total", "price per unit"):
            assert tq.product_ratio_hint(q), q
        for q in ("what is the maximum salary", "average salary per department", "total for East across Q1..Q4"):
            assert not tq.product_ratio_hint(q), q
        tq.check_spec_for_question("what is the total for East across Q1..Q4?", {"row_total": {"columns": "Q1..Q4"}})
        tq.check_spec_for_question("stock × unit price", {"derive": [{"as": "V", "expr": {}}], "row_total": {}})

    def test_agent_retries_row_total_product_with_derive(self, tmp_path):
        doc, sheets = _sheet_doc(make_inventory_xlsx(tmp_path / "inventory.xlsx"))
        prompts, kw = [], []
        bad = {"row_total": {"columns": ["Stock", "Unit price"], "as": "Total Value"},
               "aggregate": [{"column": "Total Value", "fn": "sum"}]}
        good = {"derive": [{"as": "Value", "expr": {"op": "mul", "left": "Stock", "right": "Unit price"}}],
                "aggregate": [{"column": "Value", "fn": "sum"}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(bad), json.dumps(good)], prompts, kw)
        events = _run_h(hooks, "what is the total stock value (stock × unit price)?")
        tqt = _tiles(events)["table-query"]
        assert tqt["status"] == "done"
        assert "use derive for products/ratios" in prompts[1]
        assert "| 2765 |" in prompts[-1]
        assert "1693.65" not in prompts[-1]

    def test_spec_prompt_documents_derive(self, tmp_path):
        doc, sheets = _sheet_doc(make_inventory_xlsx(tmp_path / "inventory.xlsx"))
        prompts, kw = [], []
        good = {"aggregate": [{"column": "Stock", "fn": "sum"}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(good)], prompts, kw)
        _run_h(hooks, "what is the total stock?")
        assert '"derive"' in prompts[0]
        assert "row_total\" ONLY adds columns" in prompts[0]
        assert "use \"derive\" for multiplication/division" in prompts[0].replace("Use", "use")


# ── 3. hidden sheets ────────────────────────────────────────────────────────

class TestHiddenSheets:
    def test_query_never_selects_hidden_sheet(self, company_book):
        with pytest.raises(tq.QueryError) as ei:
            tq.execute(company_book, {"sheet": "Secret", "select": ["Column A"]})
        msg = str(ei.value)
        assert "Sheet 'Secret' is hidden and not available" in msg
        assert "Available sheets: Sales, Staff" in msg and "7741" not in msg
        with pytest.raises(tq.QueryError, match="is hidden"):
            tq.execute(company_book, {"sheet": "secret"})
        with pytest.raises(tq.QueryError) as ei:
            tq.execute(company_book, {"sheet": "Budget"})
        assert str(ei.value) == "Unknown sheet 'Budget'. Available sheets: Sales, Staff"
        only_hidden = [s for s in company_book if s.get("hidden")]
        with pytest.raises(tq.QueryError, match="no visible sheets"):
            tq.execute(only_hidden, {"aggregate": ["count"]})

    def test_single_cell_sheet_is_data_not_header(self, company_book):
        secret = next(s for s in company_book if s["name"] == "Secret")
        assert list(secret["df"].columns) == ["Column A"]
        assert secret["header_row"] is None
        assert len(secret["df"]) == 1
        # a 1-column sheet with a header followed by rows keeps its header
        one_col = _sheets.build_sheet(_sheets.RawSheet("L", [["Name"], ["Ann"], ["Bob"]]))
        assert list(one_col["df"].columns) == ["Name"] and len(one_col["df"]) == 2
        lone = _sheets.build_sheet(_sheets.RawSheet("L", [["just a value"]]))
        assert list(lone["df"].columns) == ["Column A"] and lone["df"].iloc[0, 0] == "just a value"

    def test_meta_has_no_hidden_columns_or_samples(self, company_book):
        meta = _sheets.build_meta(company_book, "xlsx")
        hidden = [s for s in meta["sheets"] if s.get("hidden")]
        assert hidden == [{"name": "Secret", "hidden": True, "rows": 1, "cols": 1}]
        assert "7741" not in json.dumps(meta)
        assert meta["sheet_count"] == 2

    def test_schema_text_and_overview_count_only(self, company_book):
        for text in (tq.schema_text(company_book), _agent._sheet_overview(company_book, 5),
                     _agent._sheet_fallback_content(company_book)):
            assert "Hidden sheet(s): 1 (not shown)" in text
            assert "Secret" not in text and "7741" not in text

    def test_preview_route_hides_hidden_sheets(self, company):
        doc = asyncio.run(_idp.ingest_file("company.xlsx", company.read_bytes()))
        try:
            assert "7741" not in json.dumps(doc.meta)
            prev = _idp.sheet_preview(doc.id)
            assert [s["name"] for s in prev["sheets"]] == ["Sales", "Staff"]
            assert prev["hidden_sheets"] == 1
            assert "7741" not in json.dumps(prev) and "Secret" not in json.dumps(prev)
            with pytest.raises(LookupError, match="sheet is hidden"):
                _idp.sheet_preview(doc.id, "Secret")
            with pytest.raises(LookupError, match="sheet is hidden"):
                _idp.sheet_preview(doc.id, "secret")
            served = _idp.sheet_preview(doc.id, "Secret", include_hidden=True)
            assert served["sheet"] == "Secret" and served["rows"] == [[SECRET]]
            with pytest.raises(LookupError, match="Sheet not found"):
                _idp.sheet_preview(doc.id, "Nope")
        finally:
            _idp.delete_document(doc.id)

    def test_preview_http_route_include_hidden_param(self, company):
        """GET /api/idp/documents/{id}/sheets takes include_hidden (default False)
        and forwards it to sheet_preview. main.py is not imported (startup side
        effects): the handler is lifted out of its AST and run against _idp."""
        import ast
        import inspect
        from fastapi import HTTPException

        src = (Path(__file__).parent.parent / "main.py").read_text()
        fn = next(n for n in ast.walk(ast.parse(src))
                  if isinstance(n, ast.AsyncFunctionDef) and n.name == "idp_sheet_preview")
        route = fn.decorator_list[0]
        assert isinstance(route, ast.Call) and route.func.attr == "get"
        assert route.args[0].value == "/api/idp/documents/{doc_id}/sheets"
        fn.decorator_list = []
        ns = {"asyncio": asyncio, "_idp": _idp, "HTTPException": HTTPException}
        exec(compile(ast.Module(body=[fn], type_ignores=[]), "main.py", "exec"), ns)
        handler = ns["idp_sheet_preview"]
        param = inspect.signature(handler).parameters["include_hidden"]
        assert param.default is False and param.annotation in (bool, "bool")

        doc = asyncio.run(_idp.ingest_file("company.xlsx", company.read_bytes()))
        try:
            prev = asyncio.run(handler(doc.id))
            assert [s["name"] for s in prev["sheets"]] == ["Sales", "Staff"]
            assert prev["hidden_sheets"] == 1
            with pytest.raises(HTTPException) as exc:
                asyncio.run(handler(doc.id, "Secret"))
            assert exc.value.status_code == 404 and exc.value.detail == "sheet is hidden"
            served = asyncio.run(handler(doc.id, "Secret", include_hidden=True))
            assert served["sheet"] == "Secret" and served["rows"] == [[SECRET]]
            listed = asyncio.run(handler(doc.id, include_hidden=True))
            assert [s["name"] for s in listed["sheets"]] == ["Sales", "Staff", "Secret"]
        finally:
            _idp.delete_document(doc.id)

    def test_question_about_hidden_sheet_gets_note_not_content(self, company):
        doc, sheets = _sheet_doc(company)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [], prompts, kw)
        events = _run_h(hooks, "what is in the Secret sheet?")
        assert "table-query" not in _tiles(events)
        answer_prompt = prompts[-1]
        assert "is HIDDEN" in answer_prompt and "'Secret'" in answer_prompt
        assert "do not claim the sheet does not exist" in answer_prompt
        for pr in prompts:
            assert "7741" not in pr

    def test_count_question_about_hidden_sheet_skips_query(self, company):
        doc, sheets = _sheet_doc(company)
        prompts, kw = [], []
        spec = {"sheet": "Secret", "aggregate": ["count"]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(spec)] * 2, prompts, kw)
        events = _run_h(hooks, "how many rows and which columns are in the Secret sheet?")
        assert "table-query" not in _tiles(events)
        assert "is HIDDEN" in prompts[-1]
        for ev in events:
            assert "7741" not in json.dumps(ev, default=str)
        for pr in prompts:
            assert "7741" not in pr

    def test_mixed_question_queries_visible_sheet_and_notes_hidden(self, company):
        doc, sheets = _sheet_doc(company)
        prompts, kw = [], []
        spec = {"sheet": "Staff", "aggregate": [{"column": "Salary", "fn": "max"}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(spec)], prompts, kw)
        events = _run_h(hooks, "what is the maximum Salary in sheet Staff and list the column names of sheet Secret")
        tiles = _tiles(events)
        assert tiles["table-query"]["status"] == "done"
        assert "| 105000 |" in prompts[-1]
        assert "is HIDDEN" in prompts[-1]
        assert "Hidden sheet(s): 1 (not shown)" in prompts[0]
        for pr in prompts:
            assert "7741" not in pr

    def test_hidden_sheet_error_in_retry_prompt_has_no_content(self, company):
        doc, sheets = _sheet_doc(company)
        prompts, kw = [], []
        bad = {"sheet": "Secret", "select": ["Column A"]}
        good = {"sheet": "Staff", "aggregate": ["count"]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(bad), json.dumps(good)], prompts, kw)
        events = _run_h(hooks, "how many people are on the staff?")
        assert _tiles(events)["table-query"]["status"] == "done"
        assert "is hidden and not available" in prompts[1]
        for pr in prompts:
            assert "7741" not in pr


# ── 4. follow-ups ───────────────────────────────────────────────────────────

class TestFollowUps:
    def test_previous_question_reaches_spec_prompt(self, company):
        doc, sheets = _sheet_doc(company)
        prompts, kw = [], []
        spec = {"sheet": "Staff", "group_by": ["Department"], "aggregate": [{"column": "Salary", "fn": "max"}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(spec)], prompts, kw)
        history = [
            {"role": "user", "content": "what is the average salary per department?"},
            {"role": "assistant", "content": "Engineering 90,000; Finance 84,000; HR 57,000; Sales 62,000. " * 20},
        ]
        events = _run_h(hooks, "and the maximum?", history=history)
        assert _tiles(events)["table-query"]["status"] == "done"
        spec_prompt = prompts[0]
        assert "CONVERSATION CONTEXT" in spec_prompt
        assert "Previous question: what is the average salary per department?" in spec_prompt
        assert "resolve references like 'and the maximum?'" in spec_prompt.lower().replace("Resolve", "resolve")
        assert "QUESTION: and the maximum?" in spec_prompt
        # assistant turns are shortened
        ctx = spec_prompt.split("CONVERSATION CONTEXT")[1].split("SHEETS:")[0]
        assert len(ctx) < 1200
        assert "| Engineering | 105000 |" in prompts[-1]

    def test_context_helper(self):
        hist = [{"role": "user", "content": "q0"}, {"role": "assistant", "content": "a0"},
                {"role": "user", "content": "q1"}, {"role": "assistant", "content": "a1"},
                {"role": "user", "content": "q2"}, {"role": "assistant", "content": "a2",
                                                     "table_query": {"aggregate": [{"column": "Salary", "fn": "mean"}]}},
                {"role": "user", "content": "and the maximum?"}]
        ctx = _agent._table_conversation_context(hist, "and the maximum?")
        # the last 2 exchanges only; the current message is not its own context
        assert "q0" not in ctx and "user: q1" in ctx and "user: q2" in ctx and "assistant: a2" in ctx
        assert "and the maximum?" not in ctx.split("Resolve")[0]
        assert "Previous question: q2" in ctx
        assert "previous table query" in ctx and "Salary" in ctx
        assert _agent._table_conversation_context([], "x") == ""

    def test_no_history_no_context(self, company):
        doc, sheets = _sheet_doc(company)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps({"aggregate": [{"column": "Q2", "fn": "sum"}]})], prompts, kw)
        _run_h(hooks, "what is the total sales in Q2?")
        assert "CONVERSATION CONTEXT" not in prompts[0]


# ── 5. row_total reuses an existing Total column ───────────────────────────

class TestRowTotalReuse:
    def test_existing_formula_total_is_reused(self, company_book):
        spec = {"filters": [{"column": "Region", "op": "==", "value": "East"}],
                "row_total": {"columns": ["Q1", "Q2", "Q3", "Q4"], "as": "Total"},
                "select": ["Total"]}
        r = tq.execute(company_book, spec)
        assert r["columns"] == ["Total"]
        assert r["rows"] == [[820]]
        assert r["spec"]["row_total"]["existing"] is True
        assert "Total (computed)" not in r["markdown"]
        assert "existing column" in tq.spec_summary(r["spec"])

    def test_non_matching_total_gets_computed_name(self, company_book):
        r = tq.execute(company_book, {"row_total": {"columns": ["Q1", "Q2"], "as": "Total"},
                                      "select": ["Region"], "sort": [{"column": "Total (computed)", "desc": True}],
                                      "limit": 1})
        assert r["columns"] == ["Region", "Total (computed)"]
        assert r["rows"] == [["East", 410]]

    def test_no_total_column_keeps_requested_name(self):
        book = [_sheets.build_sheet(_sheets.RawSheet("S", [["R", "Q1", "Q2"], ["a", 1, 2], ["b", 3, 4]]))]
        r = tq.execute(book, {"row_total": {"columns": "Q1..Q2", "as": "Total"}, "select": ["R"]})
        assert r["columns"] == ["R", "Total"] and r["rows"] == [["a", 3], ["b", 7]]


# ── 6. tile detail ──────────────────────────────────────────────────────────

class TestTileDetail:
    def test_big_csv_detail(self, tmp_path):
        cities = ["Amsterdam", "Rotterdam", "Utrecht", "Eindhoven", "Groningen"]
        big = pd.DataFrame({"id": range(1, 1001)})
        big["city"] = [cities[(i - 1) % 5] for i in big.id]
        big["amount"] = big.id * 3 % 97
        p = tmp_path / "big.csv"
        big.to_csv(p, index=False)
        doc, sheets = _sheet_doc(p)
        prompts, kw = [], []
        spec = {"group_by": ["city"], "aggregate": [{"column": "amount", "fn": "sum"}]}
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [json.dumps(spec)], prompts, kw)
        events = _run_h(hooks, "what is the total amount per city?")
        tiles = _tiles(events)
        label = "Used table schema + computed result (995 rows not sent)"
        assert tiles["answer"]["detail"] == label
        assert tiles[f"extract-{doc.id}"]["detail"].endswith(label)
        assert "| Amsterdam | 9555 |" in prompts[-1]

    def test_large_sheet_without_query_detail(self, tmp_path):
        p = tmp_path / "big.csv"
        pd.DataFrame({"id": range(1, 501), "v": range(501, 1001)}).to_csv(p, index=False)
        doc, sheets = _sheet_doc(p)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [], prompts, kw)
        events = _run_h(hooks, "Summarize this spreadsheet")
        assert _tiles(events)["answer"]["detail"] == "Used table schema + stats (480 rows not sent)"


# ── 7. legacy meta ──────────────────────────────────────────────────────────

class TestLegacyMeta:
    def test_normalize_meta_sheets(self):
        meta = {"sheets": {"Sales": {"rows": 4, "cols": 6}, "Staff": {"rows": 12, "cols": 4}}}
        lst, changed = _sheets.normalize_meta_sheets(meta)
        assert changed and lst == [{"name": "Sales", "rows": 4, "cols": 6, "hidden": False},
                                   {"name": "Staff", "rows": 12, "cols": 4, "hidden": False}]
        assert _sheets.normalize_meta(meta) is True
        assert meta["sheet_count"] == 2 and meta["table_rows"] == 16
        assert _sheets.normalize_meta(meta) is False
        leaky = {"sheets": [{"name": "Secret", "hidden": True, "rows": 1, "cols": 1,
                             "columns": [{"name": SECRET, "sample": []}]}]}
        assert _sheets.normalize_meta(leaky) is True
        assert "7741" not in json.dumps(leaky)

    def test_doc_from_dict_normalises(self):
        d = {"id": "x", "filename": "a.xlsx", "mime": "application/vnd.ms-excel", "size": 1, "uploaded_at": 1.0,
             "meta": {"sheets": {"Sales": {"rows": 4, "cols": 6}}}}
        doc = _idp._doc_from_dict(d)
        assert isinstance(doc.meta["sheets"], list) and doc.meta["sheets"][0]["name"] == "Sales"

    def test_sheet_frames_refreshes_and_persists_legacy_meta(self, company):
        doc = asyncio.run(_idp.ingest_file("company.xlsx", company.read_bytes()))
        try:
            doc.meta["sheets"] = {"Sales": {"rows": 4, "cols": 6}, "Staff": {"rows": 12, "cols": 4},
                                  "Secret": {"rows": 1, "cols": 1}}
            frames = _idp.sheet_frames(doc)
            assert frames
            assert isinstance(doc.meta["sheets"], list)
            assert [s["name"] for s in doc.meta["sheets"]] == ["Sales", "Staff", "Secret"]
            assert doc.meta["sheets"][2] == {"name": "Secret", "hidden": True, "rows": 1, "cols": 1}
            assert "columns" in doc.meta["sheets"][0]
            saved = json.loads(_idp.REGISTRY_FILE.read_text())
            assert isinstance(saved[doc.id]["meta"]["sheets"], list)
            assert "7741" not in json.dumps(saved[doc.id]["meta"])
        finally:
            _idp.delete_document(doc.id)

    def test_agent_tolerates_dict_meta(self, company):
        doc, sheets = _sheet_doc(company)
        doc.meta = {"sheets": {"Sales": {"rows": 4, "cols": 6}, "Staff": {"rows": 12, "cols": 4}}}
        assert _agent.is_sheet_doc(doc)
        prompts, kw = [], []
        hooks = _hooks({doc.id: doc}, {doc.id: sheets}, [], prompts, kw)
        events = _run_h(hooks, "Summarize this spreadsheet")
        assert _tiles(events)[f"extract-{doc.id}"]["detail"] == "Spreadsheet · 2 sheets · 16 rows"
